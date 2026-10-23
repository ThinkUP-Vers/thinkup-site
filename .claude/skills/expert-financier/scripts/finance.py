#!/usr/bin/env python3
"""expert-financier : données, analyse quantitative, backtests, simulations, note de stratégie, journal.

Aucun chiffre ne se calcule à la main : statistiques, tests hors échantillon,
allocations par budget, frais, simulations de Monte-Carlo, probabilités du
modèle statistique et scores du journal sont calculés ici, à partir de
analyse.json (voir references/analyse-json.md) et des séries de prix.

  finance.py init [racine] --instrument "LVMH" [--date AAAA-MM-JJ]
  finance.py donnees <dossier> --symbole MC.PA [--source yahoo|stooq|fred] [--nom mc] [--debut 2000-01-01]
  finance.py importer <dossier> <fichier.csv> --nom mc [--colonne Close] [--source "..."]
  finance.py analyser <dossier> [--aujourdhui AAAA-MM-JJ]
  finance.py backtest <dossier> [--fichier donnees/x.csv] [--signaux s.csv] [--journal DIR]
  finance.py construire <dossier> [--final] [--journal DIR] [--chemins N]
  finance.py obligation --prix 98.5 --coupon 3 --echeance 2034-05-25 [--frequence 1]
  finance.py evaluer [--journal DIR] [--donnees DIR ...] [--prix SYMBOLE|ID=VALEUR ...]
  finance.py annoter <id> --cause analyse|evenement_exogene|timing|hasard|donnees [--note "..."]
  finance.py bilan [--journal DIR] [--sortie DIR]
  finance.py archiver [racine]          finance.py restaurer <archive.zip> [racine]

--aujourdhui AAAA-MM-JJ fixe la date d'arrêté : les séries sont coupées à cette
date (analyse « à date », rejouable sur le passé). Par défaut : date du jour.
--journal vaut par défaut finance/journal.

Code de sortie : 0 si aucune erreur, 1 sinon. Les alertes n'empêchent rien
mais doivent être lues.
"""
import argparse
import copy
import csv
import datetime as dt
import html
import importlib.util
import io
import json
import math
import re
import shutil
import subprocess
import sys
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

try:
    import numpy as np
    from numpy.lib.stride_tricks import sliding_window_view
except ImportError:  # pragma: no cover
    sys.exit("numpy est requis : pip install numpy")

SKILL = Path(__file__).resolve().parent.parent
HORIZONS = ("court", "moyen", "long")
LIBELLE_HORIZON = {"court": "Court terme", "moyen": "Moyen terme", "long": "Long terme"}
JOURS_BOURSE = 252
CAUSES = ("analyse", "evenement_exogene", "timing", "hasard", "donnees")
PROBABILITES_EVT = ("faible", "moyenne", "elevee", "certaine")
# Repères S&P 500 (pic, creux) ; approximatifs pour les autres marchés.
FENETRES_STRESS = [
    ("Crise financière 2008", "2007-10-09", "2009-03-09"),
    ("Crise de la zone euro 2011", "2011-04-29", "2011-10-03"),
    ("Resserrement monétaire fin 2018", "2018-09-20", "2018-12-24"),
    ("Krach Covid 2020", "2020-02-19", "2020-03-23"),
    ("Choc des taux 2022", "2022-01-03", "2022-10-12"),
    ("Choc tarifaire 2025", "2025-02-19", "2025-04-08"),
]
CSS = """
body{font-family:Georgia,'Times New Roman',serif;color:#1d1d1f;background:#f4f4f2;margin:0}
.feuille{max-width:980px;margin:24px auto;background:#fff;padding:40px 48px;box-shadow:0 1px 4px rgba(0,0,0,.08)}
h1{font-size:1.6em;margin:.2em 0 .4em}h2{font-size:1.25em;border-bottom:1px solid #ddd;padding-bottom:.2em;margin-top:1.6em}
h3{font-size:1.05em;margin-top:1.2em}table{border-collapse:collapse;width:100%;margin:.6em 0 1em;font-size:.88em}
th,td{border:1px solid #ccc;padding:4px 6px;text-align:left;vertical-align:top}th{background:#f0f0ee}
blockquote{border-left:4px solid #b54;margin:1em 0;padding:.4em 1em;background:#fbf3f1}
@media (max-width:640px){.feuille{padding:20px 16px}table{display:block;overflow-x:auto}}
"""


# ── Outils généraux ──────────────────────────────────────────────────────

def sortir(msg):
    print(f"ERREUR : {msg}", file=sys.stderr)
    sys.exit(1)


class Rapport:
    def __init__(self):
        self.erreurs, self.alertes = [], []

    def erreur(self, msg):
        self.erreurs.append(msg)

    def alerte(self, msg):
        self.alertes.append(msg)

    def afficher(self, titre="Contrôle"):
        print(f"\n== {titre} : {len(self.erreurs)} erreur(s), {len(self.alertes)} alerte(s)")
        for m in self.erreurs:
            print(f"  ERREUR  {m}")
        for m in self.alertes:
            print(f"  alerte  {m}")


def lire_json(chemin):
    try:
        return json.loads(Path(chemin).read_text(encoding="utf-8"))
    except FileNotFoundError:
        sortir(f"{chemin} introuvable")
    except json.JSONDecodeError as e:
        sortir(f"{chemin} : JSON invalide ({e})")


def ecrire_json(chemin, donnees):
    Path(chemin).write_text(json.dumps(donnees, ensure_ascii=False, indent=2, default=_json_defaut) + "\n",
                            encoding="utf-8")


def _json_defaut(o):
    if isinstance(o, (np.floating, np.integer)):
        return o.item()
    if isinstance(o, np.ndarray):
        return o.tolist()
    if isinstance(o, (dt.date, dt.datetime)):
        return o.isoformat()
    raise TypeError(type(o))


def slug(texte):
    t = unicodedata.normalize("NFKD", texte).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")[:40] or "instrument"


def date_iso(texte, champ, R=None):
    try:
        return dt.date.fromisoformat(str(texte))
    except (TypeError, ValueError):
        if R is not None:
            R.erreur(f"{champ} : date attendue au format AAAA-MM-JJ, reçu « {texte} »")
        return None


def fusionner(base, ajout):
    out = copy.deepcopy(base)
    for k, v in (ajout or {}).items():
        out[k] = fusionner(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else v
    return out


def charger_config(analyse=None, racine=Path("finance")):
    cfg = lire_json(SKILL / "config.json")
    local = Path(racine) / "config.local.json"
    if local.exists():
        cfg = fusionner(cfg, lire_json(local))
    if analyse:
        cfg = fusionner(cfg, analyse.get("parametres") or {})
    return cfg


def fini(x):
    return x is not None and isinstance(x, (int, float, np.floating)) and math.isfinite(float(x))


def nb(x, dec=2):
    if not fini(x):
        return "n.d."
    return f"{x:,.{dec}f}".replace(",", " ").replace(".", ",").replace(" ", " ")


def eur(x, dec=0):
    return "n.d." if not fini(x) else f"{nb(x, dec)} €"


def pct(x, dec=1, signe=False):
    if not fini(x):
        return "n.d."
    s = f"{x * 100:+.{dec}f}" if signe else f"{x * 100:.{dec}f}"
    return s.replace(".", ",") + " %"


def tableau(tete, lignes):
    nettoie = lambda c: str(c).replace("|", "/").replace("\n", " ")
    out = ["| " + " | ".join(nettoie(c) for c in tete) + " |", "|" + "---|" * len(tete)]
    out += ["| " + " | ".join(nettoie(c) for c in l) + " |" for l in lignes]
    return "\n".join(out)


def aujourdhui(a):
    return dt.date.fromisoformat(a.aujourdhui) if getattr(a, "aujourdhui", None) else dt.date.today()


def duree(jours):
    if jours < 365:
        return f"{jours} jours"
    ans = round(jours / 365.25, 1)
    return "1 an" if ans == 1 else f"{nb(ans, 0 if ans == int(ans) else 1)} ans"


def jours_bourse(jours_calendaires):
    return max(1, round(jours_calendaires * JOURS_BOURSE / 365))


# ── Séries de prix ───────────────────────────────────────────────────────

FORMATS_DATE = ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y%m%d", "%d-%m-%Y", "%Y/%m/%d",
                "%b %d, %Y", "%d %b %Y", "%Y-%m-%d %H:%M:%S"]
COLONNES_PRIX = ["adj close", "adjclose", "adjusted close", "cours ajuste", "cloture ajustee", "cloture",
                 "close", "clot", "clot.", "dernier", "price", "last", "vl", "nav", "valeur", "value", "cours"]
COLONNES_DATE = ["date", "datetime", "time", "jour", "observation_date", "timestamp"]


def sans_accent(t):
    return unicodedata.normalize("NFKD", str(t)).encode("ascii", "ignore").decode().strip().lower()


def lire_serie(chemin, jusqua=None):
    """Lit une série normalisée (date,cloture) ; coupe à la date d'arrêté."""
    dates, prix = [], []
    with open(chemin, encoding="utf-8") as f:
        lecteur = csv.reader(f)
        tete = next(lecteur, None)
        if not tete or sans_accent(tete[0]) != "date":
            sortir(f"{chemin} : série non normalisée (en-tête date,cloture attendu) ; passer par « importer »")
        for ligne in lecteur:
            if len(ligne) < 2 or not ligne[1].strip():
                continue
            d = dt.date.fromisoformat(ligne[0])
            if jusqua and d > jusqua:
                break
            if dates and d <= dates[-1]:
                sortir(f"{chemin} : dates non croissantes vers {d} ; repasser le fichier par « importer »")
            dates.append(d)
            prix.append(float(ligne[1]))
    if len(prix) < 2:
        sortir(f"{chemin} : moins de deux cours avant le {jusqua}")
    return dates, np.array(prix, dtype=float)


def nombre(texte):
    t = str(texte).strip().strip('"').replace(" ", "").replace(" ", "").replace(" ", "")
    t = t.replace("€", "").replace("$", "").replace("%", "")
    if t in ("", ".", "-", "null", "None", "nan", "N/A", "n/a"):
        return None
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    try:
        v = float(t)
    except ValueError:
        return None
    return v if math.isfinite(v) else None


def normaliser_csv(texte, colonne=None, positif=True):
    """Lit un CSV d'origine quelconque ; renvoie ([(date, valeur)] croissant sans doublon, colonne retenue)."""
    texte = texte.lstrip("﻿")
    echantillon = texte[:4096]
    try:
        dialecte = csv.Sniffer().sniff(echantillon, delimiters=",;\t|")
    except csv.Error:
        dialecte = csv.excel
    lignes = [l for l in csv.reader(io.StringIO(texte), dialecte) if any(c.strip() for c in l)]
    if len(lignes) < 3:
        sortir("fichier vide ou illisible")
    tete = [sans_accent(c) for c in lignes[0]]
    i_date = next((i for i, c in enumerate(tete) if c in COLONNES_DATE), 0)
    if colonne:
        cible = sans_accent(colonne)
        if cible not in tete:
            sortir(f"colonne « {colonne} » absente ; colonnes : {', '.join(lignes[0])}")
        i_prix = tete.index(cible)
    else:
        i_prix = next((tete.index(c) for c in COLONNES_PRIX if c in tete), None)
        if i_prix is None:
            if len(tete) == 2:
                i_prix = 1 - i_date
            else:
                sortir(f"colonne de cours introuvable ; préciser --colonne parmi : {', '.join(lignes[0])}")
    brutes = [(l[i_date].strip().strip('"'), l[i_prix]) for l in lignes[1:] if len(l) > max(i_date, i_prix)]
    textes_dates = [d for d, _ in brutes]
    choix = None
    for fmt in FORMATS_DATE:
        try:
            ds = [dt.datetime.strptime(d, fmt).date() for d in textes_dates]
        except ValueError:
            continue
        monotone = all(a <= b for a, b in zip(ds, ds[1:])) or all(a >= b for a, b in zip(ds, ds[1:]))
        if monotone:
            choix = ds
            break
        choix = choix or ds
    if choix is None:
        sortir(f"format de date non reconnu (ex. « {textes_dates[0]} »)")
    serie = {}
    for d, (_, v) in zip(choix, brutes):
        x = nombre(v)
        if x is not None and (x > 0 or not positif):
            serie[d] = x
    if len(serie) < 2:
        sortir("aucun cours exploitable")
    return sorted(serie.items()), lignes[0][i_prix].strip()


def ecrire_serie(chemin, points):
    chemin.parent.mkdir(parents=True, exist_ok=True)
    with open(chemin, "w", encoding="utf-8", newline="") as f:
        f.write("date,cloture\n")
        for d, v in points:
            f.write(f"{d.isoformat()},{v:.6g}\n")


def noter_source(dossier, nom, info):
    chemin = dossier / "donnees" / "sources.json"
    sources = json.loads(chemin.read_text(encoding="utf-8")) if chemin.exists() else {}
    sources[f"donnees/{nom}.csv"] = info
    ecrire_json(chemin, sources)


def http(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (expert-financier)"})
    hote = urllib.parse.urlparse(url).hostname
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.read().decode("utf-8", "replace")
    except (urllib.error.URLError, OSError) as e:
        motif = str(getattr(e, "reason", e))
        if "403" in motif or "Tunnel" in motif:
            sortir(f"accès refusé à {hote} (politique réseau de l'environnement). Autoriser ce domaine dans "
                   f"les réglages réseau de l'environnement, ou exporter le CSV et lancer « finance.py importer ».")
        sortir(f"{hote} injoignable : {motif}")


def telecharger(source, symbole, debut):
    if source == "yahoo":
        p1 = int(dt.datetime.combine(debut, dt.time(), dt.timezone.utc).timestamp())
        p2 = int(dt.datetime.now(dt.timezone.utc).timestamp())
        url = (f"https://query1.finance.yahoo.com/v8/finance/chart/{urllib.parse.quote(symbole)}"
               f"?period1={p1}&period2={p2}&interval=1d&events=div%2Csplit&includeAdjustedClose=true")
        try:
            res = json.loads(http(url))["chart"]["result"][0]
            decalage = res["meta"].get("gmtoffset", 0) or 0
            cours = (res["indicators"].get("adjclose") or [{}])[0].get("adjclose") \
                or res["indicators"]["quote"][0]["close"]
            points = {dt.datetime.fromtimestamp(t + decalage, dt.timezone.utc).date(): float(v)
                      for t, v in zip(res["timestamp"], cours) if v}
            ajuste = bool((res["indicators"].get("adjclose") or [{}])[0].get("adjclose"))
        except (KeyError, IndexError, TypeError, json.JSONDecodeError):
            sortir(f"réponse Yahoo inattendue pour {symbole} (symbole inconnu ou format changé)")
        return sorted(points.items()), "Yahoo Finance (API chart non officielle)", ajuste
    if source == "stooq":
        url = f"https://stooq.com/q/d/l/?s={urllib.parse.quote(symbole.lower())}&i=d&d1={debut:%Y%m%d}"
        return normaliser_csv(http(url))[0], "Stooq (CSV quotidien)", None
    if source == "fred":
        url = (f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={urllib.parse.quote(symbole)}"
               f"&cosd={debut.isoformat()}")
        return normaliser_csv(http(url), positif=False)[0], "FRED, Federal Reserve Bank of St. Louis", None
    sortir(f"source inconnue : {source}")


# ── Statistiques ─────────────────────────────────────────────────────────

def moyenne_mobile(p, n):
    out = np.full(len(p), np.nan)
    if len(p) >= n:
        c = np.cumsum(np.insert(p, 0, 0.0))
        out[n - 1:] = (c[n:] - c[:-n]) / n
    return out


def vol_glissante(p, n=20):
    out = np.full(len(p), np.nan)
    r = p[1:] / p[:-1] - 1
    if len(r) >= n:
        out[n:] = sliding_window_view(r, n).std(axis=1, ddof=1) * math.sqrt(JOURS_BOURSE)
    return out


def rsi(p, n=14):
    out = np.full(len(p), np.nan)
    d = np.diff(p)
    if len(d) < n:
        return out
    g, l = np.maximum(d, 0), np.maximum(-d, 0)
    mg, ml = g[:n].mean(), l[:n].mean()
    out[n] = 100.0 if ml == 0 else 100 - 100 / (1 + mg / ml)
    for i in range(n, len(d)):
        mg = (mg * (n - 1) + g[i]) / n
        ml = (ml * (n - 1) + l[i]) / n
        out[i + 1] = 100.0 if ml == 0 else 100 - 100 / (1 + mg / ml)
    return out


def mesures(r, rf):
    """Mesures d'une série de rendements quotidiens simples."""
    r = np.asarray(r, dtype=float)
    if len(r) < 20:
        return None
    v = np.concatenate([[1.0], np.cumprod(1 + r)])
    ans = len(r) / JOURS_BOURSE
    vol = r.std(ddof=1) * math.sqrt(JOURS_BOURSE)
    exces = r - ((1 + rf) ** (1 / JOURS_BOURSE) - 1)
    mdd = float((v / np.maximum.accumulate(v) - 1).min())
    cagr = v[-1] ** (1 / ans) - 1 if v[-1] > 0 else -1.0
    return {"rendement_annuel": cagr, "volatilite": vol,
            "sharpe": exces.mean() * JOURS_BOURSE / vol if vol > 0 else None,
            "perte_max": mdd, "calmar": cagr / abs(mdd) if mdd < 0 else None, "annees": ans}


def perte_max_detail(dates, p):
    pic = np.maximum.accumulate(p)
    dd = p / pic - 1
    i_creux = int(dd.argmin())
    i_pic = int(p[:i_creux + 1].argmax())
    apres = np.nonzero(p[i_creux:] >= p[i_pic])[0]
    i_recup = i_creux + int(apres[0]) if len(apres) else None
    return {"perte_max": float(dd[i_creux]), "pic": dates[i_pic], "creux": dates[i_creux],
            "recuperation": dates[i_recup] if i_recup is not None else None,
            "jours_sous_l_eau": (dates[i_recup] - dates[i_pic]).days if i_recup is not None
            else (dates[-1] - dates[i_pic]).days}


def cours_a(dates, p, jour):
    """Dernier cours connu au plus tard le jour donné (ou None)."""
    i = int(np.searchsorted(np.array(dates, dtype="datetime64[D]"), np.datetime64(jour), side="right")) - 1
    return (dates[i], float(p[i])) if i >= 0 else (None, None)


def stress(dates, p):
    out = []
    for nom, a, b in FENETRES_STRESS:
        da, db = dt.date.fromisoformat(a), dt.date.fromisoformat(b)
        if dates[0] > da or dates[-1] < db:
            out.append({"episode": nom, "debut": a, "fin": b, "variation": None, "perte_max": None})
            continue
        i = int(np.searchsorted(np.array(dates, dtype="datetime64[D]"), np.datetime64(da)))
        j = int(np.searchsorted(np.array(dates, dtype="datetime64[D]"), np.datetime64(db), side="right")) - 1
        seg = p[i:j + 1]
        out.append({"episode": nom, "debut": a, "fin": b, "variation": float(seg[-1] / seg[0] - 1),
                    "perte_max": float((seg / np.maximum.accumulate(seg) - 1).min())})
    return out


def analyser_serie(dates, p, rf, ref=None):
    r = p[1:] / p[:-1] - 1
    m = mesures(r, rf) or {}
    q05 = np.quantile(r, 0.05)
    s = r.std(ddof=1)
    z = (r - r.mean()) / s if s > 0 else np.zeros_like(r)
    perf = {}
    for nom, k in (("1_mois", 21), ("3_mois", 63), ("6_mois", 126), ("1_an", 252), ("3_ans", 756), ("5_ans", 1260)):
        perf[nom] = float(p[-1] / p[-1 - k] - 1) if len(p) > k else None
    glissant = lambda k: float((p[k:] / p[:-k] - 1).min()) if len(p) > k else None
    mm50, mm200 = moyenne_mobile(p, 50)[-1], moyenne_mobile(p, 200)[-1]
    vols = vol_glissante(p, 20)
    vols_ok = vols[~np.isnan(vols)]
    vol20 = float(vols[-1]) if len(vols_ok) else None
    atr = float(np.abs(r[-14:]).mean()) if len(r) >= 14 else None
    an = p[-252:] if len(p) >= 252 else p
    res = {
        "debut": dates[0], "fin": dates[-1], "dernier_cours": float(p[-1]), "observations": len(p),
        **m,
        "perte_max_detail": perte_max_detail(dates, p),
        "var_95_jour": float(-q05), "cvar_95_jour": float(-r[r <= q05].mean()),
        "pire_jour": float(r.min()), "pire_mois": glissant(21), "pire_annee": glissant(252),
        "asymetrie": float((z ** 3).mean()), "aplatissement": float((z ** 4).mean() - 3),
        "performances": perf,
        "momentum_12_1": float(p[-22] / p[-253] - 1) if len(p) > 253 else None,
        "mm50": float(mm50) if not math.isnan(mm50) else None,
        "mm200": float(mm200) if not math.isnan(mm200) else None,
        "rsi14": float(rsi(p, 14)[-1]) if len(p) > 15 else None,
        "volatilite_20j": vol20,
        "percentile_vol_20j": float((vols_ok <= vol20).mean()) if vol20 is not None else None,
        "variation_moyenne_14j": atr,
        "stops_suggeres": {"2_fois": 2 * atr, "3_fois": 3 * atr} if atr else None,
        "plus_haut_52s": float(an.max()), "plus_bas_52s": float(an.min()),
        "stress": stress(dates, p),
    }
    if ref is not None:
        dref, pref = ref
        communs = sorted(set(dates) & set(dref))
        if len(communs) > 60:
            ia = {d: i for i, d in enumerate(dates)}
            ib = {d: i for i, d in enumerate(dref)}
            a = np.array([p[ia[d]] for d in communs])
            b = np.array([pref[ib[d]] for d in communs])
            ra, rb = a[1:] / a[:-1] - 1, b[1:] / b[:-1] - 1
            res["beta"] = float(np.cov(ra, rb)[0, 1] / rb.var(ddof=1))
            res["correlation"] = float(np.corrcoef(ra, rb)[0, 1])
    return res


# ── Stratégies et tests hors échantillon ─────────────────────────────────

def s_achat(p):
    return np.ones(len(p))


def s_tendance(p, n):
    m = moyenne_mobile(p, n)
    return np.where(np.isnan(m), 0.0, (p > m).astype(float))


def s_croisement(p, rapide, lente):
    a, b = moyenne_mobile(p, rapide), moyenne_mobile(p, lente)
    return np.where(np.isnan(b), 0.0, (a > b).astype(float))


def s_momentum(p, retard):
    e = np.zeros(len(p))
    if len(p) > retard:
        e[retard:] = (p[retard - 21:len(p) - 21] / p[:len(p) - retard] > 1).astype(float)
    return e


def s_vol_cible(p, cible, bande=0.1):
    """Exposition cible / volatilité 20 jours, ajustée seulement quand l'écart dépasse la bande (frais)."""
    v = vol_glissante(p, 20)
    voulu = np.where(np.isnan(v) | (v <= 0), 0.0, np.minimum(1.0, cible / np.where(v > 0, v, 1)))
    e = np.zeros(len(p))
    for t in range(1, len(p)):
        e[t] = voulu[t] if abs(voulu[t] - e[t - 1]) > bande else e[t - 1]
    return e


def s_rsi_court(p, seuil):
    """RSI(2) sous le seuil dans une tendance haussière (cours > MM200) ; sortie au-dessus de la MM5."""
    r2, m200, m5 = rsi(p, 2), moyenne_mobile(p, 200), moyenne_mobile(p, 5)
    e, dedans = np.zeros(len(p)), False
    for t in range(len(p)):
        if np.isnan(m200[t]) or np.isnan(r2[t]):
            continue
        if not dedans and r2[t] < seuil and p[t] > m200[t]:
            dedans = True
        elif dedans and p[t] > m5[t]:
            dedans = False
        e[t] = 1.0 if dedans else 0.0
    return e


def s_donchian(p, entree, sortie):
    e, dedans = np.zeros(len(p)), False
    for t in range(max(entree, sortie), len(p)):
        if not dedans and p[t] > p[t - entree:t].max():
            dedans = True
        elif dedans and p[t] < p[t - sortie:t].min():
            dedans = False
        e[t] = 1.0 if dedans else 0.0
    return e


STRATEGIES = {
    "achat_conservation": {"fn": s_achat, "grille": [{}], "horizons": ["moyen", "long"],
                           "libelle": "Achat et conservation"},
    "tendance_mm": {"fn": s_tendance, "grille": [{"n": 100}, {"n": 150}, {"n": 200}], "horizons": ["moyen", "long"],
                    "libelle": "Suivi de tendance : investi si le cours dépasse sa moyenne mobile"},
    "croisement_mm": {"fn": s_croisement, "grille": [{"rapide": 20, "lente": 100}, {"rapide": 50, "lente": 200}],
                      "horizons": ["moyen"], "libelle": "Croisement de moyennes mobiles"},
    "momentum_12_1": {"fn": s_momentum, "grille": [{"retard": 126}, {"retard": 252}], "horizons": ["moyen"],
                      "libelle": "Momentum absolu (performance passée hors dernier mois)"},
    "volatilite_cible": {"fn": s_vol_cible, "grille": [{"cible": 0.10}, {"cible": 0.15}, {"cible": 0.20}],
                         "horizons": ["moyen", "long"], "libelle": "Exposition réduite quand la volatilité monte"},
    "rsi_court": {"fn": s_rsi_court, "grille": [{"seuil": 5}, {"seuil": 10}, {"seuil": 20}], "horizons": ["court"],
                  "libelle": "Achat sur repli court (RSI 2) dans une tendance haussière"},
    "cassure_donchian": {"fn": s_donchian, "grille": [{"entree": 20, "sortie": 10}, {"entree": 55, "sortie": 20}],
                         "horizons": ["court", "moyen"], "libelle": "Cassure de plus haut (canal de Donchian)"},
}


def charger_modeles(dossiers, R):
    """Modèles ajoutés : fichiers .py définissant NOM, LIBELLE, HORIZONS, GRILLE et expositions(prix, **p)."""
    modeles = {}
    for d in dossiers:
        for f in sorted(Path(d).glob("*.py")) if Path(d).is_dir() else []:
            spec = importlib.util.spec_from_file_location(f"modele_{f.stem}", f)
            mod = importlib.util.module_from_spec(spec)
            try:
                spec.loader.exec_module(mod)
                if mod.NOM in STRATEGIES or mod.NOM in modeles:
                    R.erreur(f"modèle {f} : le nom {mod.NOM} est déjà pris, modèle ignoré")
                    continue
                modeles[mod.NOM] = {"fn": mod.expositions, "grille": list(mod.GRILLE) or [{}],
                                    "horizons": list(mod.HORIZONS), "libelle": mod.LIBELLE, "fichier": str(f)}
            except Exception as e:  # un modèle cassé ne doit pas bloquer les autres
                R.erreur(f"modèle {f} illisible : {e}")
    return modeles


def sans_anticipation(fn, p, params):
    """Vrai si l'exposition du jour t ne dépend que des cours jusqu'à t (contrôlé à cinq coupures)."""
    complet = np.asarray(fn(p, **params), dtype=float)
    if complet.shape != p.shape:
        return False
    for coupe in sorted({len(p) // 4, len(p) // 2, 3 * len(p) // 4, len(p) - 30, len(p) - 2}):
        if coupe < 30:
            continue
        partiel = np.asarray(fn(p[:coupe + 1], **params), dtype=float)
        if not np.allclose(np.nan_to_num(partiel), np.nan_to_num(complet[:coupe + 1]), atol=1e-9):
            return False
    return True


def frais_ordre(montant, courtier):
    """Frais d'un ordre (vectorisé) : courtage et écart de cotation."""
    m = np.asarray(montant, dtype=float)
    courtage = np.maximum(courtier["frais_min"], courtier["frais_fixes"] + courtier["frais_pct"] * m)
    return np.where(m > 0, courtage + courtier["spread"] * m, 0.0)


def rendements_strategie(p, e, rf, courtier, budget):
    r = p[1:] / p[:-1] - 1
    pos = np.clip(np.nan_to_num(np.asarray(e, dtype=float)), 0, 1)[:-1]
    rf_j = (1 + rf) ** (1 / JOURS_BOURSE) - 1
    brut = pos * r + (1 - pos) * rf_j
    variation = np.abs(np.diff(np.concatenate([[0.0], pos])))
    return brut - frais_ordre(variation * budget, courtier) / budget


def walk_forward(p, strat, rf, courtier, budget, entrainement, test):
    series = [rendements_strategie(p, strat["fn"](p, **g), rf, courtier, budget) for g in strat["grille"]]
    expos = [np.clip(np.nan_to_num(np.asarray(strat["fn"](p, **g), dtype=float)), 0, 1) for g in strat["grille"]]
    hors, choix, sharpes_in, expo_hors = [], [], [], []
    debut = entrainement
    while debut + test <= len(p) - 1:
        notes = []
        for s in series:
            m = mesures(s[debut - entrainement:debut], rf)
            notes.append(m["sharpe"] if m and m["sharpe"] is not None else -9.0)
        k = int(np.argmax(notes))
        hors.append(series[k][debut:debut + test])
        expo_hors.append(expos[k][debut:debut + test])
        choix.append(strat["grille"][k])
        sharpes_in.append(notes[k])
        debut += test
    if not hors:
        return None
    r = np.concatenate(hors)
    e = np.concatenate(expo_hors)
    m = mesures(r, rf)
    m.update({"temps_investi": float(e.mean()),
              "transactions_par_an": float((np.abs(np.diff(e)) > 1e-9).sum() / (len(e) / JOURS_BOURSE)),
              "sharpe_entrainement": float(np.mean(sharpes_in)),
              "parametres_retenus": max(({json.dumps(c, sort_keys=True): choix.count(c) for c in choix}).items(),
                                        key=lambda kv: kv[1])[0],
              "fenetres": len(hors), "debut_idx": entrainement})
    return m


# ── Obligations ──────────────────────────────────────────────────────────

def ajouter_mois(d, mois):
    a, m = divmod(d.month - 1 + mois, 12)
    an, mo = d.year + a, m + 1
    jours_mois = [31, 29 if an % 4 == 0 and (an % 100 or an % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    return dt.date(an, mo, min(d.day, jours_mois[mo - 1]))


def obligation(prix, coupon, echeance, reglement, frequence=1, nominal=100.0):
    """Rendement actuariel, duration et convexité d'une obligation à taux fixe (prix pied de coupon en % du nominal)."""
    pas = 12 // frequence
    dates, d = [], echeance
    while d > reglement:
        dates.append(d)
        d = ajouter_mois(echeance, -pas * len(dates))
    precedent = d
    dates.reverse()
    c = coupon / frequence * nominal / 100
    couru = c * (reglement - precedent).days / max(1, (dates[0] - precedent).days)
    sale = prix * nominal / 100 + couru
    t = np.array([(x - reglement).days / 365.25 for x in dates])
    flux = np.full(len(dates), c)
    flux[-1] += nominal
    f = frequence
    valeur = lambda y: float((flux / (1 + y / f) ** (f * t)).sum())
    bas, haut = -0.5, 1.0
    for _ in range(200):
        milieu = (bas + haut) / 2
        bas, haut = (milieu, haut) if valeur(milieu) > sale else (bas, milieu)
    y = (bas + haut) / 2
    va = flux / (1 + y / f) ** (f * t)
    macaulay = float((t * va).sum() / sale)
    modifiee = macaulay / (1 + y / f)
    convexite = float((flux * t * (t + 1 / f) / (1 + y / f) ** (f * t + 2)).sum() / sale)
    choc = lambda dy: -modifiee * dy + 0.5 * convexite * dy ** 2
    return {"rendement_actuariel": y, "coupon_couru": couru, "prix_plein_coupon": sale, "duration_macaulay": macaulay,
            "duration_modifiee": modifiee, "convexite": convexite, "variation_si_taux_plus_1pt": choc(0.01),
            "variation_si_taux_moins_1pt": choc(-0.01), "flux_restants": len(dates)}


# ── Allocation, frais, fiscalité ─────────────────────────────────────────

def frais_achat(montant, ligne, cfg):
    f = float(frais_ordre(montant, cfg["courtier"]))
    if ligne.get("ttf") and montant > 0:
        f += cfg["ttf"]["taux"] * montant
    return f


def allouer(lignes, budget, cfg):
    sortie, total = [], 0.0
    for l in lignes:
        cible = l["poids"] * budget
        prix = l.get("prix")
        if prix and not l.get("fractionnable"):
            q = math.floor(cible / prix)
            while q > 0 and q * prix + frais_achat(q * prix, l, cfg) > cible + 1e-9:
                q -= 1
            montant = q * prix
        else:
            q, montant = None, cible
            for _ in range(8):
                montant = max(0.0, cible - frais_achat(montant, l, cfg))
        f = frais_achat(montant, l, cfg) if montant > 0 else 0.0
        sortie.append({"nom": l["nom"], "quantite": q, "montant": montant, "frais": f,
                       "frais_pct": f / montant if montant > 0 else None})
        total += montant + f
    return sortie, budget - total


def taux_impot(enveloppe, date_sortie, cfg):
    fisc = cfg["fiscalite"]
    if enveloppe == "PEA":
        ouvert = fisc.get("pea_ouvert_le")
        if not ouvert:
            return fisc["PEA"]["avant_5_ans"], "PEA, date d'ouverture inconnue : taux avant 5 ans retenu"
        age = (date_sortie - dt.date.fromisoformat(ouvert)).days / 365.25
        return (fisc["PEA"]["apres_5_ans"], "PEA de plus de 5 ans à la sortie") if age >= 5 else \
            (fisc["PEA"]["avant_5_ans"], "PEA de moins de 5 ans à la sortie")
    if enveloppe == "AV":
        ouverte = fisc.get("av_ouverte_le")
        if ouverte and (date_sortie - dt.date.fromisoformat(ouverte)).days / 365.25 >= 8:
            return fisc["AV"]["apres_8_ans"], "assurance-vie de plus de 8 ans, gain sous l'abattement"
        return fisc["AV"]["avant_8_ans"], "assurance-vie de moins de 8 ans (ou date d'ouverture inconnue)"
    if enveloppe == "SOCIETE":
        return fisc["SOCIETE"]["taux"], "compte-titres d'une société à l'IS (hors distribution à l'associé)"
    return fisc["CTO"]["taux"], "compte-titres ordinaire (PFU)"


# ── Simulations de Monte-Carlo ───────────────────────────────────────────

def tirages(T, H, n, bloc, graine):
    rng = np.random.default_rng(graine)
    bloc = min(bloc, T)
    nb_blocs = -(-H // bloc)
    debuts = rng.integers(0, T - bloc + 1, size=(n, nb_blocs))
    return (debuts[:, :, None] + np.arange(bloc)).reshape(n, nb_blocs * bloc)[:, :H]


def fenetres_historiques(T, H, pas=5):
    """Toutes les fenêtres réelles de H séances (une tous les « pas » jours) : rejoue les crises dans leur ordre."""
    if T < H:
        return None
    return np.arange(0, T - H + 1, pas)[:, None] + np.arange(H)


def bloc_effectif(H, cfg):
    """Blocs plus longs pour les longs horizons, afin de garder la durée des crises (2008 a duré 17 mois)."""
    return max(cfg["simulation"]["bloc_jours"], min(126, H // 10))


CLASSE_PAR_TYPE = {"action": "actions", "obligation": "obligations", "matiere_premiere": "matieres_premieres",
                   "devise": "devises", "crypto": "crypto", "liquidites": "monetaire", "produit_levier": "monetaire"}
MODES = ("central", "neutre", "historique")
ENVELOPPES = ("PEA", "CTO", "AV", "SOCIETE")


def classe_de(element, cfg):
    """Classe d'actif d'une ligne : explicite, sinon déduite du type (None si indéterminable : ETF, fonds, indice)."""
    c = element.get("classe") or CLASSE_PAR_TYPE.get(element.get("type"))
    return c if c in cfg["primes_de_risque"] else None


def ajuster_derive(lr, mode, rf, primes=None):
    """Recentre la tendance de chaque colonne : en valeur médiane, chaque actif croît au taux sans risque
    (neutre), au taux sans risque plus la prime prudente de sa classe (central), ou garde sa tendance passée."""
    if mode == "historique":
        return lr
    primes = np.zeros(lr.shape[1]) if mode == "neutre" or primes is None else np.asarray(primes, dtype=float)
    return lr - lr.mean(axis=0) + np.log1p(rf + primes) / JOURS_BOURSE


def appliquer_levier(lr, leviers):
    """Produit à levier quotidien constant (approximation) : rendement simple multiplié, perte bornée à la mise."""
    if all(l == 1 for l in leviers):
        return lr
    r = np.expm1(lr) * np.asarray(leviers, dtype=float)
    return np.log1p(np.maximum(r, -0.9999))


def simuler(lr, idx, lignes, alloc, liquidites, budget, cfg, rf, reequilibrage, lr_ref=None):
    """Valeur d'un portefeuille le long de chemins rééchantillonnés par blocs, stops et plancher compris."""
    n, H = idx.shape
    m0 = np.array([a["montant"] for a in alloc], dtype=float)
    stop = np.array([l.get("stop") or np.inf for l in lignes], dtype=float)
    objectif = np.array([l.get("objectif") or np.inf for l in lignes], dtype=float)
    suiveur = np.array([bool(l.get("stop_suiveur")) for l in lignes])
    poids_cible = np.array([l["poids"] for l in lignes], dtype=float)
    val = np.tile(m0, (n, 1))
    pic_ligne = val.copy()
    ouvert = np.tile(m0 > 0, (n, 1))
    cash = np.full(n, float(liquidites))
    plancher = budget * (1 - cfg["seuil_perte_max"])
    touche = np.zeros(n, bool)
    pic = np.full(n, float(budget))
    mdd = np.zeros(n)
    croissance_cash = math.exp(math.log1p(rf) / JOURS_BOURSE)
    courtier = cfg["courtier"]
    ref = None
    if lr_ref is not None:
        ref = np.full(n, budget - float(frais_ordre(budget, courtier)))
    seuil_objectif = np.where(np.isinf(objectif), np.inf, m0 * (1 + np.where(np.isinf(objectif), 0, objectif)))
    facteur_stop = np.where(np.isinf(stop), -np.inf, 1 - np.where(np.isinf(stop), 0, stop))
    for t in range(H):
        val = np.where(ouvert, val * np.exp(lr[idx[:, t]]), 0.0)
        cash = cash * croissance_cash
        if ref is not None:
            ref = ref * np.exp(lr_ref[idx[:, t]])
        pic_ligne = np.maximum(pic_ligne, val)
        base = np.where(suiveur, pic_ligne, m0)
        sortie = ouvert & ((val <= base * np.where(base > 0, facteur_stop, 0)) | (val >= seuil_objectif))
        if sortie.any():
            cash = cash + np.where(sortie, val - frais_ordre(val, courtier), 0.0).sum(axis=1)
            val = np.where(sortie, 0.0, val)
            ouvert = ouvert & ~sortie
        if reequilibrage and (t + 1) % JOURS_BOURSE == 0 and t + 1 < H:
            # Lignes infaisables pour ce budget (montant nul) hors du rééquilibrage : leur part reste en liquidités.
            complet = (ouvert | (m0 <= 0)).all(axis=1) & ~touche
            if complet.any():
                V = val.sum(axis=1) + cash
                cible = V[:, None] * np.where(m0 > 0, poids_cible, 0.0)[None, :]
                ecart = np.abs(cible - val)
                cout = (frais_ordre(ecart, courtier) * (ecart > 0.01 * V[:, None])).sum(axis=1)
                val = np.where(complet[:, None], cible, val)
                cash = np.where(complet, V - cible.sum(axis=1) - cout, cash)
        V = val.sum(axis=1) + cash
        nouveau = ~touche & (V <= plancher)
        if nouveau.any():
            cash = np.where(nouveau, cash + (val - frais_ordre(val, courtier)).sum(axis=1), cash)
            val = np.where(nouveau[:, None], 0.0, val)
            ouvert = ouvert & ~nouveau[:, None]
            touche |= nouveau
            V = val.sum(axis=1) + cash
        pic = np.maximum(pic, V)
        mdd = np.minimum(mdd, V / pic - 1)
    final = val.sum(axis=1) - frais_ordre(val, courtier).sum(axis=1) + cash
    ans = H / JOURS_BOURSE
    q = np.quantile(final, [0.05, 0.25, 0.5, 0.75, 0.95])
    res = {"p05": q[0], "p25": q[1], "mediane": q[2], "p75": q[3], "p95": q[4], "moyenne": float(final.mean()),
           "pire_valeur": float(final.min()),
           "proba_perte": float((final < budget).mean()), "proba_plancher": float(touche.mean()),
           "perte_max_mediane": float(np.median(mdd)), "perte_max_p95": float(np.quantile(mdd, 0.05)),
           "rendement_annualise_median": float((q[2] / budget) ** (1 / ans) - 1) if q[2] > 0 else -1.0,
           "proba_sortie_anticipee": float((~ouvert & ~touche[:, None] & (m0 > 0)).any(axis=1).mean())}
    if ref is not None:
        ref_final = ref - frais_ordre(ref, courtier)
        res.update({"reference_mediane": float(np.median(ref_final)),
                    "reference_p05": float(np.quantile(ref_final, 0.05)),
                    "proba_bat_reference": float((final > ref_final).mean())})
    return res


def proba_modele(lr, H, ratio, sens, cfg, rf, prime=0.0):
    """Probabilité que le cours finisse au-dessus (ou en dessous) du seuil, si sa tendance médiane est le taux sans
    risque plus la prime prudente de sa classe et son risque celui de son historique."""
    s = cfg["simulation"]
    lr_n = ajuster_derive(lr.reshape(-1, 1), "central", rf, [prime])[:, 0]
    idx = tirages(len(lr_n), H, max(s["chemins"], 4000), bloc_effectif(H, cfg), s["graine"] + 1)
    final = lr_n[idx].sum(axis=1)
    seuil = math.log(ratio)
    return float((final > seuil).mean()) if sens == "au_dessus" else float((final < seuil).mean())


# ── Lecture et contrôle d'analyse.json ───────────────────────────────────

def chemin_donnee(dossier, rel):
    return (dossier / rel) if rel else None


def charger_series(dossier, analyse, jusqua, R):
    """Toutes les séries citées par analyse.json, coupées à la date d'arrêté."""
    citees = {}
    inst = analyse.get("instrument", {})
    ref = analyse.get("reference_passive") or {}
    for rel in [inst.get("fichier"), ref.get("fichier")] + [
            l.get("proxy") or l.get("fichier") for h in (analyse.get("strategies") or {}).values()
            for l in (h or {}).get("lignes", [])] + [p.get("fichier") for p in analyse.get("previsions", [])]:
        if rel and rel not in citees:
            c = chemin_donnee(dossier, rel)
            if not c.exists():
                R.erreur(f"série {rel} introuvable dans {dossier}")
                continue
            citees[rel] = lire_serie(c, jusqua)
    return citees


def aligner(series, fichiers):
    communs = None
    for f in fichiers:
        ds = set(series[f][0])
        communs = ds if communs is None else communs & ds
    communs = sorted(communs or [])
    mat = []
    for f in fichiers:
        dates, p = series[f]
        pos = {d: i for i, d in enumerate(dates)}
        mat.append(np.array([p[pos[d]] for d in communs]))
    return communs, np.array(mat).T if mat else None


def controler(analyse, cfg, series, asof, R, journal):
    meta = analyse.get("meta", {})
    for champ in ("reference", "demande"):
        if not meta.get(champ):
            R.erreur(f"meta.{champ} manquant")
    inst = analyse.get("instrument", {})
    for champ in ("nom", "type", "fichier"):
        if not inst.get(champ):
            R.erreur(f"instrument.{champ} manquant")
    if inst.get("type") in cfg["types_interdits"]:
        R.erreur(f"instrument de type {inst['type']} : perte possible au-delà de la mise, interdit par les règles")
    if inst.get("prix") is not None:
        if not inst.get("prix_source") or not inst.get("prix_date"):
            R.erreur("instrument.prix sans prix_source ni prix_date : un cours se source et se date")
        else:
            d = date_iso(inst["prix_date"], "instrument.prix_date", R)
            if d and (asof - d).days > 5:
                R.alerte(f"cours de l'instrument daté du {d} ({(asof - d).days} jours) : le relever à nouveau")
    if inst.get("fichier") in series:
        dates = series[inst["fichier"]][0]
        if (asof - dates[-1]).days > 5:
            R.alerte(f"série {inst['fichier']} arrêtée au {dates[-1]} : mettre les données à jour")
        if (dates[-1] - dates[0]).days < 365 * cfg["simulation"]["historique_min_annees"]:
            R.alerte(f"historique de {inst['fichier']} inférieur à {cfg['simulation']['historique_min_annees']} ans : "
                     "statistiques et simulations peu fiables")
    if inst.get("devise") and inst["devise"] != "EUR":
        R.alerte(f"instrument coté en {inst['devise']} : le risque de change n'est compté que si la série est "
                 "convertie en euros")
    rep = analyse.get("repartition") or {}
    for b, parts in rep.items():
        if b not in {str(x) for x in cfg["budgets"]}:
            R.erreur(f"repartition.{b} : budget absent de la configuration ({', '.join(map(str, cfg['budgets']))})")
        if abs(sum(parts.values()) - 1) > 0.001:
            R.erreur(f"repartition.{b} : les parts doivent sommer à 1")
        for h_ in parts:
            if h_ not in (analyse.get("strategies") or {}):
                R.erreur(f"repartition.{b} : horizon {h_} sans stratégie")
    if not (analyse.get("reference_passive") or {}).get("fichier"):
        R.alerte("aucune référence passive : impossible de dire si la stratégie bat un simple ETF")
    an = analyse.get("analyse", {})
    if not an.get("these"):
        R.erreur("analyse.these manquante")
    if len(an.get("premortem") or []) < 2:
        R.erreur("analyse.premortem : au moins deux raisons pour lesquelles la thèse pourrait échouer")
    for i, f in enumerate(an.get("fondamental") or []):
        if not f.get("source") or not f.get("date"):
            R.erreur(f"analyse.fondamental[{i}] sans source ni date : « {str(f.get('point'))[:60]} »")
    for i, v in enumerate(an.get("valorisation") or []):
        if not v.get("source") or not v.get("date"):
            R.erreur(f"analyse.valorisation[{i}] ({v.get('indicateur')}) sans source ni date")
    evts = analyse.get("evenements") or []
    if not evts:
        R.erreur("aucun événement exogène recensé : la veille est obligatoire (references/exogenes.md)")
    for i, e in enumerate(evts):
        if not e.get("source") or not e.get("verifie_le"):
            R.erreur(f"evenements[{i}] ({e.get('evenement')}) sans source ni date de vérification")
        else:
            d = date_iso(e["verifie_le"], f"evenements[{i}].verifie_le", R)
            if d and (asof - d).days > 7:
                R.alerte(f"événement « {e.get('evenement')} » vérifié il y a {(asof - d).days} jours")
        if e.get("probabilite") not in PROBABILITES_EVT:
            R.erreur(f"evenements[{i}].probabilite : {'|'.join(PROBABILITES_EVT)}")
        if not e.get("parade"):
            R.erreur(f"evenements[{i}] ({e.get('evenement')}) sans parade décidée à l'avance")
    sc = analyse.get("scenarios") or []
    if sc and abs(sum(s.get("probabilite", 0) for s in sc) - 1) > 0.01:
        R.erreur("scenarios : les probabilités doivent sommer à 1")
    strategies = analyse.get("strategies") or {}
    if not strategies:
        R.erreur("aucune stratégie")
    for h, s in strategies.items():
        if h not in HORIZONS:
            R.erreur(f"strategies.{h} : horizon inconnu ({'|'.join(HORIZONS)})")
            continue
        lignes = s.get("lignes") or []
        somme = sum(l.get("poids", 0) for l in lignes) + s.get("liquidites", 0)
        if abs(somme - 1) > 0.001:
            R.erreur(f"strategies.{h} : poids + liquidités = {somme:.3f}, 1 attendu")
        if not s.get("invalidation"):
            R.erreur(f"strategies.{h}.invalidation manquante : ce qui prouverait que la thèse est fausse")
        env = s.get("enveloppe") or (analyse.get("parametres") or {}).get("enveloppe", "CTO")
        if env not in ENVELOPPES:
            R.erreur(f"strategies.{h} : enveloppe « {env} » inconnue ({', '.join(ENVELOPPES)})")
        if env == "PEA":
            for l in lignes:
                if l.get("type") in ("obligation", "crypto", "matiere_premiere", "devise") or \
                        classe_de(l, cfg) in ("obligations", "matieres_premieres", "crypto", "devises"):
                    R.alerte(f"{h} / {l.get('nom')} : a priori non éligible au PEA (actions européennes et ETF "
                             "éligibles seulement) ; changer d'enveloppe ou de support")
        for l in lignes:
            nom = l.get("nom", "?")
            typ = l.get("type", inst.get("type"))
            if not l.get("nom") or not fini(l.get("poids")) or not 0 <= l["poids"] <= 1:
                R.erreur(f"{h} / {nom} : nom et poids (fraction entre 0 et 1) obligatoires")
            for cle in ("stop", "objectif"):
                if l.get(cle) is not None and (not fini(l[cle]) or not 0 < l[cle] < (1 if cle == "stop" else 10)):
                    R.erreur(f"{h} / {nom} : {cle} en fraction attendu (0,08 pour 8 %), reçu {l[cle]}")
            if typ in cfg["types_interdits"] or l.get("perte_illimitee"):
                R.erreur(f"{h} / {nom} : produit à perte potentiellement supérieure à la mise, interdit")
            elif typ not in cfg["types_autorises"]:
                R.erreur(f"{h} / {nom} : type « {typ} » inconnu")
            if l.get("classe") and l["classe"] not in cfg["primes_de_risque"]:
                R.erreur(f"{h} / {nom} : classe « {l['classe']} » inconnue ({', '.join(k for k in cfg['primes_de_risque'] if not k.startswith('_'))})")
            elif typ in ("etf", "fonds", "indice") and not classe_de(l, cfg):
                R.erreur(f"{h} / {nom} : classe d'actif obligatoire pour un {typ} (actions, obligations...), "
                         "elle fixe la prime de risque du cas central")
            if typ == "produit_levier":
                if not fini(l.get("levier")) or not 1 < l["levier"] <= 20:
                    R.erreur(f"{h} / {nom} : levier (entre 1 et 20) obligatoire pour un produit à levier")
                if h != "court":
                    R.erreur(f"{h} / {nom} : produit à levier réservé au court terme")
                if l.get("poids", 0) > cfg["poids_max_levier"]:
                    R.erreur(f"{h} / {nom} : produit à levier au-delà de {pct(cfg['poids_max_levier'], 0)} du budget")
            if not (l.get("proxy") or l.get("fichier")):
                R.erreur(f"{h} / {nom} : ni fichier ni proxy, simulation impossible")
            if l.get("proxy"):
                R.alerte(f"{h} / {nom} : simulé avec le proxy {l['proxy']}")
            plafond = cfg["risque_max_par_ligne"].get(h)
            if h in ("court", "moyen") and not l.get("stop"):
                R.erreur(f"{h} / {nom} : stop obligatoire à cet horizon")
            if plafond and l.get("stop") and l.get("poids", 0) * l["stop"] > plafond + 1e-9:
                R.erreur(f"{h} / {nom} : risque {pct(l['poids'] * l['stop'])} du budget au stop, "
                         f"plafond {pct(plafond)} ; réduire le poids ou resserrer le stop")
            if h == "court" and l.get("stop") and l.get("objectif") and l["objectif"] / l["stop"] < 1.5:
                R.alerte(f"court / {nom} : gain visé {pct(l['objectif'])} pour {pct(l['stop'])} de risque "
                         f"(ratio {nb(l['objectif'] / l['stop'], 1)}, 1,5 au moins conseillé)")
            if typ == "action" and l.get("poids", 0) > 0.5 and h == "long":
                R.alerte(f"long / {nom} : plus de la moitié du budget sur une seule action")
    prevs = analyse.get("previsions") or []
    for h in strategies:
        if not any(p.get("horizon") == h for p in prevs):
            R.erreur(f"aucune prévision chiffrée pour l'horizon {h} : sans elle, la stratégie n'est pas évaluable")
    ids = [p.get("id") for p in prevs]
    if len(ids) != len(set(ids)):
        R.erreur("previsions : identifiants en double")
    for p in prevs:
        nom = p.get("id", "?")
        if not p.get("enonce"):
            R.erreur(f"prévision {nom} : énoncé manquant")
        cond = p.get("condition") or {}
        if cond.get("type") not in ("au_dessus", "en_dessous") or not fini(cond.get("seuil")):
            R.erreur(f"prévision {nom} : condition {{type: au_dessus|en_dessous, seuil}} attendue")
        if not fini(p.get("probabilite")) or not 0.01 <= p["probabilite"] <= 0.99:
            R.erreur(f"prévision {nom} : probabilité entre 0,01 et 0,99 attendue")
        e = date_iso(p.get("echeance"), f"prévision {nom}.echeance", R)
        if e and e <= asof:
            R.erreur(f"prévision {nom} : échéance {e} déjà passée")
        if not p.get("fichier"):
            R.erreur(f"prévision {nom} : fichier de cours manquant (sert à l'évaluation)")
    lecons = Path(journal) / "lecons.md"
    if lecons.exists() and not analyse.get("lecons_appliquees"):
        R.alerte(f"{lecons} existe mais analyse.lecons_appliquees est vide : dire quelles leçons ont joué")


# ── Construction de la note ──────────────────────────────────────────────

def calculer(dossier, analyse, cfg, asof, R, journal, chemins=None):
    series = charger_series(dossier, analyse, asof, R)
    rf = cfg["taux_sans_risque"]["valeur"]
    inst = analyse.get("instrument", {})
    ref = analyse.get("reference_passive") or {}
    calc = {"date_arrete": asof, "instrument": None, "reference": None, "strategies": {}, "previsions": {},
            "obligation": None}
    if inst.get("fichier") in series:
        calc["instrument"] = analyser_serie(*series[inst["fichier"]], rf,
                                            series.get(ref.get("fichier")) if ref.get("fichier") else None)
    if ref.get("fichier") in series:
        calc["reference"] = analyser_serie(*series[ref["fichier"]], rf)
    if inst.get("obligation"):
        o = inst["obligation"]
        try:
            calc["obligation"] = obligation(o["prix"], o["coupon"], dt.date.fromisoformat(o["echeance"]), asof,
                                            o.get("frequence", 1))
        except (KeyError, ValueError) as e:
            R.erreur(f"instrument.obligation incomplet : {e}")
    controler(analyse, cfg, series, asof, R, journal)
    s_cfg = cfg["simulation"]
    n = chemins or s_cfg["chemins"]
    for h, s in (analyse.get("strategies") or {}).items():
        lignes = s.get("lignes") or []
        fichiers = [l.get("proxy") or l.get("fichier") for l in lignes]
        if not lignes or any(f not in series for f in fichiers) or any(
                not l.get("nom") or not fini(l.get("poids")) or not 0 <= l["poids"] <= 1
                or (l.get("stop") is not None and not (fini(l["stop"]) and 0 < l["stop"] < 1)) for l in lignes):
            continue
        avec_ref = ref.get("fichier") in series
        dates, mat = aligner(series, fichiers + ([ref["fichier"]] if avec_ref else []))
        if mat is None or len(dates) < 260:
            R.erreur(f"strategies.{h} : moins d'un an d'historique commun aux lignes, simulation impossible")
            continue
        if (dates[-1] - dates[0]).days < 365 * s_cfg["historique_min_annees"]:
            R.alerte(f"strategies.{h} : historique commun de {(dates[-1] - dates[0]).days // 365} an(s) seulement")
        lr_tout = np.log(mat[1:] / mat[:-1])
        primes_cfg = cfg["primes_de_risque"]
        primes = [primes_cfg.get(classe_de(l, cfg), 0.0) for l in lignes] + \
            ([primes_cfg.get(classe_de(ref, cfg) or "actions", 0.0)] if avec_ref else [])
        leviers = [float(l.get("levier") or 1) if l.get("type") == "produit_levier" else 1.0 for l in lignes] + \
            ([1.0] if avec_ref else [])
        jours = s.get("horizon_jours") or cfg["horizons_jours"][h]
        H = jours_bourse(jours)
        bloc = bloc_effectif(H, cfg)
        idx = tirages(len(lr_tout), H, n, bloc, s_cfg["graine"])
        rejeu = fenetres_historiques(len(lr_tout), H)
        sortie = {"horizon_jours": jours, "jours_bourse": H, "historique": [dates[0], dates[-1]], "budgets": {},
                  "primes": {l["nom"]: p for l, p in zip(lignes, primes)},
                  "bloc_jours": bloc, "fenetres_rejouees": 0 if rejeu is None else len(rejeu)}
        date_sortie = asof + dt.timedelta(days=jours)
        taux, motif = taux_impot(s.get("enveloppe") or analyse.get("parametres", {}).get("enveloppe", "CTO"),
                                 date_sortie, cfg)
        sortie["impot"] = {"taux": taux, "motif": motif}
        for budget in cfg["budgets"]:
            alloc, liquidites = allouer(lignes, budget, cfg)
            for a, l in zip(alloc, lignes):
                if l["poids"] > 0 and a["montant"] <= 0:
                    R.alerte(f"{h} / {l['nom']} : impossible avec {eur(budget)} (cours supérieur à la part allouée)")
                elif a["frais_pct"] and a["frais_pct"] > 0.01:
                    R.alerte(f"{h} / {l['nom']} avec {eur(budget)} : frais d'achat de {pct(a['frais_pct'])} de l'ordre "
                             "(au-delà de 1 %, la ligne est trop petite pour ce budget)")
            res = {"allocation": alloc, "liquidites": liquidites}
            for mode in MODES:
                lr = appliquer_levier(ajuster_derive(lr_tout, mode, rf, primes), leviers)
                res[mode] = simuler(lr[:, :len(lignes)], idx, lignes, alloc, liquidites, budget, cfg, rf,
                                    bool(s.get("reequilibrage_annuel")), lr[:, -1] if avec_ref else None)
            if rejeu is not None:
                lr = appliquer_levier(lr_tout, leviers)
                res["rejeu"] = simuler(lr[:, :len(lignes)], rejeu, lignes, alloc, liquidites, budget, cfg, rf,
                                       bool(s.get("reequilibrage_annuel")), lr[:, -1] if avec_ref else None)
            pire = max([res[m]["proba_plancher"] for m in MODES] + [res.get("rejeu", {}).get("proba_plancher", 0)])
            res["proba_plancher_max"] = pire
            if pire > cfg["proba_max_plancher"]:
                R.erreur(f"{h} avec {eur(budget)} : probabilité {pct(pire)} de perdre "
                         f"{pct(cfg['seuil_perte_max'], 0)} de la mise, plafond {pct(cfg['proba_max_plancher'], 0)}")
            gain = res["central"]["mediane"] - budget
            res["gain_median_net_impot"] = gain * (1 - taux) if gain > 0 else gain
            sortie["budgets"][str(budget)] = res
        calc["strategies"][h] = sortie
    for p in analyse.get("previsions") or []:
        f = p.get("fichier")
        if f not in series or not (p.get("condition") or {}).get("seuil"):
            continue
        dates, prix = series[f]
        e = date_iso(p.get("echeance"), "echeance")
        if not e or e <= asof:
            continue
        prix_ref = p.get("prix_reference") or float(prix[-1])
        lr = np.log(prix[1:] / prix[:-1])
        classe = classe_prevision(p, analyse, cfg)
        if classe is None:
            R.alerte(f"prévision {p['id']} : classe d'actif inconnue (champ classe), modèle sans prime de risque")
        prime = cfg["primes_de_risque"].get(classe, 0.0)
        pm = proba_modele(lr, jours_bourse((e - asof).days), p["condition"]["seuil"] / prix_ref,
                          p["condition"]["type"], cfg, rf, prime)
        calc["previsions"][p["id"]] = {"probabilite_modele": pm, "prix_reference": prix_ref, "classe": classe,
                                       "ecart": p["probabilite"] - pm}
        if abs(p["probabilite"] - pm) >= 0.25:
            R.alerte(f"prévision {p['id']} : {pct(p['probabilite'], 0)} contre {pct(pm, 0)} pour le modèle "
                     "statistique ; un tel écart exige un argument précis et sourcé dans la thèse")
    return calc, series


def classe_prevision(p, analyse, cfg):
    """Classe d'actif de la série d'une prévision : explicite, sinon celle de la ligne, de l'instrument ou de la
    référence qui partage son fichier."""
    if p.get("classe") in cfg["primes_de_risque"]:
        return p["classe"]
    f = p.get("fichier")
    candidats = [l for s in (analyse.get("strategies") or {}).values() for l in (s or {}).get("lignes", [])
                 if (l.get("proxy") or l.get("fichier")) == f]
    candidats += [x for x in (analyse.get("instrument") or {}, analyse.get("reference_passive") or {})
                  if x.get("fichier") == f]
    for c in candidats:
        if classe_de(c, cfg):
            return classe_de(c, cfg)
    return "actions" if (analyse.get("reference_passive") or {}).get("fichier") == f else None


def ligne_budget(budget, res):
    n = res["central"]
    return [eur(budget), eur(n["mediane"]), f"{eur(n['p05'])} à {eur(n['p95'])}", pct(n["proba_perte"], 0),
            pct(res["proba_plancher_max"], 1), eur(n.get("reference_mediane")), pct(n.get("proba_bat_reference"), 0)]


def note_md(analyse, calc, cfg, R, backtest, sources):
    inst = analyse.get("instrument", {})
    meta = analyse.get("meta", {})
    ref = analyse.get("reference_passive") or {}
    an = analyse.get("analyse", {})
    asof = calc["date_arrete"]
    out = [f"# Note d'analyse : {inst.get('nom', '?')}" + (f" ({inst['symbole']})" if inst.get("symbole") else ""),
           "",
           f"_Référence {meta.get('reference')} · données arrêtées au {asof} · usage personnel de Patrick, "
           "pas un conseil en investissement au sens réglementaire_", ""]
    if any(v.get("synthetique") for v in sources.values()):
        out += ["> DONNÉES SYNTHÉTIQUES : cette note sert d'exemple et ne décrit aucun marché réel.", ""]
    out += ["## En bref", "", f"**Thèse.** {an.get('these', '')}", ""]
    for h in HORIZONS:
        s = (analyse.get("strategies") or {}).get(h)
        c = calc["strategies"].get(h)
        if not s or not c:
            continue
        medianes = ", ".join(f"{eur(float(b))} → {eur(r['central']['mediane'])}" for b, r in c["budgets"].items())
        n = (c["budgets"].get("5000") or next(iter(c["budgets"].values())))["central"]
        out.append(f"- **{LIBELLE_HORIZON[h]} ({duree(c['horizon_jours'])}).** {s.get('objectif', '')} "
                   f"Valeurs médianes simulées (cas central) : {medianes}. Pour 5 000 € : 5 % des cas sous "
                   f"{eur(n['p05'])}, probabilité de perte {pct(n['proba_perte'], 0)}.")
    if analyse.get("recommandation"):
        out += ["", f"**Recommandation.** {analyse['recommandation']}"]
    rep = analyse.get("repartition") or {}
    if rep:
        hs = [h for h in HORIZONS if any(h in p for p in rep.values())]
        out += ["", tableau(["Budget"] + [LIBELLE_HORIZON[h] for h in hs],
                            [[eur(float(b))] + [eur(float(b) * p.get(h, 0)) if p.get(h) else "—" for h in hs]
                             for b, p in rep.items()])]
    out.append("")
    if R.alertes or R.erreurs:
        out += ["## Points d'attention", ""]
        out += [f"- **Bloquant** : {m}" for m in R.erreurs] + [f"- {m}" for m in R.alertes] + [""]

    out += ["## 1. Instrument", ""]
    lignes = [["Nom", inst.get("nom", "")], ["Type", inst.get("type", "")]]
    for cle, lib in (("symbole", "Symbole"), ("isin", "ISIN"), ("place", "Place de cotation"), ("devise", "Devise")):
        if inst.get(cle):
            lignes.append([lib, inst[cle]])
    if "eligible_pea" in inst:
        lignes.append(["Éligible au PEA", "oui" if inst["eligible_pea"] else "non"])
    if inst.get("prix") is not None:
        lignes.append(["Dernier cours relevé", f"{nb(inst['prix'])} ({inst.get('prix_date')}, {inst.get('prix_source')})"])
    out += [tableau(["", ""], lignes), ""]

    q = calc.get("instrument")
    if q:
        r_ = calc.get("reference") or {}
        pmd = q["perte_max_detail"]
        out += ["## 2. Analyse quantitative", "",
                f"Historique du {q['debut']} au {q['fin']} ({q['observations']} séances). "
                f"Taux sans risque retenu : {pct(cfg['taux_sans_risque']['valeur'])} "
                f"({cfg['taux_sans_risque']['libelle']}, vérifié le {cfg['taux_sans_risque']['verifie_le']}).", ""]
        lig = [["Rendement annualisé", pct(q.get("rendement_annuel")), pct(r_.get("rendement_annuel"))],
               ["Volatilité annualisée", pct(q.get("volatilite")), pct(r_.get("volatilite"))],
               ["Ratio de Sharpe", nb(q.get("sharpe")), nb(r_.get("sharpe"))],
               ["Perte maximale historique", f"{pct(pmd['perte_max'])} ({pmd['pic']} → {pmd['creux']})",
                pct((r_.get("perte_max_detail") or {}).get("perte_max"))],
               ["Retour au plus haut", str(pmd["recuperation"] or "pas encore") +
                f" ({pmd['jours_sous_l_eau']} jours sous le pic)", ""],
               ["Perte journalière dépassée 1 jour sur 20 (VaR 95 %)", pct(q["var_95_jour"]), pct(r_.get("var_95_jour"))],
               ["Pire mois / pire année glissants", f"{pct(q['pire_mois'])} / {pct(q['pire_annee'])}",
                f"{pct(r_.get('pire_mois'))} / {pct(r_.get('pire_annee'))}"],
               ["Performance 1 an / 3 ans", f"{pct(q['performances']['1_an'])} / {pct(q['performances']['3_ans'])}",
                f"{pct((r_.get('performances') or {}).get('1_an'))} / {pct((r_.get('performances') or {}).get('3_ans'))}"],
               ["Momentum 12 mois hors dernier mois", pct(q.get("momentum_12_1")), pct(r_.get("momentum_12_1"))]]
        out += [tableau(["Mesure", inst.get("nom", "Instrument"), ref.get("nom", "Référence")], lig), ""]
        tech = []
        if q.get("mm50") and q.get("mm200"):
            tech.append(f"cours {nb(q['dernier_cours'])} ; moyenne mobile 50 jours {nb(q['mm50'])}, "
                        f"200 jours {nb(q['mm200'])} : tendance "
                        f"{'haussière' if q['dernier_cours'] > q['mm200'] else 'baissière'} de fond")
        if q.get("rsi14") is not None:
            tech.append(f"RSI 14 jours {nb(q['rsi14'], 0)}" + (" (suracheté)" if q["rsi14"] > 70 else
                                                                " (survendu)" if q["rsi14"] < 30 else ""))
        if q.get("volatilite_20j") is not None:
            tech.append(f"volatilité sur 20 jours {pct(q['volatilite_20j'])}, plus haute que "
                        f"{pct(q['percentile_vol_20j'], 0)} de l'historique")
        tech.append(f"plus haut / plus bas sur 52 semaines : {nb(q['plus_haut_52s'])} / {nb(q['plus_bas_52s'])}")
        if q.get("beta") is not None:
            tech.append(f"bêta {nb(q['beta'])} et corrélation {nb(q['correlation'])} avec {ref.get('nom', 'la référence')}")
        if q.get("stops_suggeres"):
            tech.append(f"stop de court terme cohérent avec la volatilité (2 à 3 fois la variation quotidienne moyenne "
                        f"sur 14 jours) : {pct(q['stops_suggeres']['2_fois'])} à {pct(q['stops_suggeres']['3_fois'])}")
        out += [f"- {t}" for t in tech] + [""]
        if an.get("technique"):
            out += [an["technique"], ""]
        st = [[s_["episode"], f"{s_['debut']} → {s_['fin']}", pct(s_["variation"]), pct(s_["perte_max"]),
               pct(next((x["variation"] for x in (r_.get("stress") or []) if x["episode"] == s_["episode"]), None))]
              for s_ in q["stress"] if s_["variation"] is not None]
        if st:
            out += ["### Tests de crise (épisodes historiques)", "",
                    tableau(["Épisode", "Période", "Variation", "Pire repli", ref.get("nom", "Référence")], st), ""]

    if calc.get("obligation"):
        o = calc["obligation"]
        out += ["### Obligation", "", tableau(["Mesure", "Valeur"], [
            ["Rendement actuariel", pct(o["rendement_actuariel"], 2)],
            ["Duration modifiée", nb(o["duration_modifiee"])],
            ["Convexité", nb(o["convexite"])],
            ["Variation du prix si les taux montent d'un point", pct(o["variation_si_taux_plus_1pt"])],
            ["Variation du prix si les taux baissent d'un point", pct(o["variation_si_taux_moins_1pt"])],
            ["Coupon couru", nb(o["coupon_couru"])]]), ""]

    if an.get("fondamental") or an.get("valorisation"):
        out += ["## 3. Analyse fondamentale et valorisation", ""]
        out += [f"- {f['point']} _({f['source']}, {f['date']})_" for f in an.get("fondamental") or []]
        if an.get("valorisation"):
            out += ["", tableau(["Indicateur", "Valeur", "Comparaison", "Source"],
                                [[v.get("indicateur"), v.get("valeur"), v.get("comparaison", ""),
                                  f"{v.get('source')}, {v.get('date')}"] for v in an["valorisation"]])]
        out.append("")

    evts = analyse.get("evenements") or []
    if evts:
        out += ["## 4. Événements exogènes à surveiller", "",
                tableau(["Date", "Événement", "Probabilité", "Effet attendu", "Parade décidée à l'avance", "Source"],
                        [[e.get("date") or "non daté", e.get("evenement"), e.get("probabilite"),
                          f"{e.get('sens', '')} {e.get('ampleur_estimee', '')}".strip(), e.get("parade"),
                          f"{e.get('source')} (vérifié le {e.get('verifie_le')})"]
                         for e in sorted(evts, key=lambda e: e.get("date") or "9999")]), ""]

    sc = analyse.get("scenarios") or []
    if sc:
        esp = sum(s_.get("probabilite", 0) * s_.get("rendement", 0) for s_ in sc)
        out += ["## 5. Scénarios", "",
                tableau(["Scénario", "Probabilité", "Rendement sur l'horizon", "Ce qui le déclenche"],
                        [[s_.get("nom"), pct(s_.get("probabilite"), 0), pct(s_.get("rendement"), 0, True),
                          s_.get("description", "")] for s_ in sc]),
                "", f"Espérance pondérée : {pct(esp, 1, True)}.", ""]

    out += ["## 6. Stratégies par horizon et par budget", "",
            f"Plancher absolu : perte de {pct(cfg['seuil_perte_max'], 0)} de la mise ; une stratégie n'est retenue que si "
            f"la probabilité de l'atteindre reste sous {pct(cfg['proba_max_plancher'], 0)}. "
            f"Frais : {cfg['courtier']['nom']}.", ""]
    for h in HORIZONS:
        s = (analyse.get("strategies") or {}).get(h)
        c = calc["strategies"].get(h)
        if not s:
            continue
        out += [f"### {LIBELLE_HORIZON[h]} : {s.get('objectif', '')}", ""]
        for cle, lib in (("entree", "Entrée"), ("invalidation", "Invalidation"), ("sortie", "Sortie")):
            if s.get(cle):
                out.append(f"- **{lib}.** {s[cle]}")
        if s.get("reequilibrage_annuel"):
            out.append("- **Rééquilibrage** annuel vers les poids cibles.")
        out.append("")
        lig = []
        for l in s.get("lignes") or []:
            p = l.get("prix")
            lig.append([l.get("nom", "?"), pct(l.get("poids"), 0), nb(p) if p else "n.d.",
                        f"{nb(p * (1 - l['stop'])) if p and l.get('stop') else ''} ({pct(l['stop'], 0)})"
                        + (" suiveur" if l.get("stop_suiveur") else "") if l.get("stop") else "aucun",
                        f"{nb(p * (1 + l['objectif'])) if p else ''} ({pct(l['objectif'], 0, True)})"
                        if l.get("objectif") else "aucun",
                        nb(l["objectif"] / l["stop"], 1) if l.get("objectif") and l.get("stop") else ""])
        if s.get("liquidites"):
            lig.append(["Liquidités", pct(s["liquidites"], 0), "", "", "", ""])
        out += [tableau(["Ligne", "Poids", "Cours de référence", "Stop", "Objectif", "Gain / risque"], lig), ""]
        if not c:
            out += ["_Simulation impossible (voir points d'attention)._", ""]
            continue
        alloc_l = []
        for b, r in c["budgets"].items():
            qte = ", ".join(f"{a['nom']} : {a['quantite'] if a['quantite'] is not None else eur(a['montant'])}"
                            for a in r["allocation"])
            fr = sum(a["frais"] for a in r["allocation"])
            alloc_l.append([eur(float(b)), qte, eur(sum(a["montant"] for a in r["allocation"])),
                            f"{eur(fr, 2)} ({pct(fr / float(b), 2)})", eur(r["liquidites"])])
        out += [tableau(["Budget", "Quantités", "Investi", "Frais d'achat (part du budget)", "Liquidités"], alloc_l), ""]
        out += [f"Résultats simulés à {duree(c['horizon_jours'])} ({c['jours_bourse']} séances), "
                f"{cfg['simulation']['chemins']} chemins tirés par blocs de {c['bloc_jours']} séances "
                f"dans l'historique du {c['historique'][0]} au {c['historique'][1]}. Cas central : risque de "
                f"l'historique, tendance médiane égale au taux sans risque plus une prime prudente par classe d'actif ("
                + ", ".join(f"{k} {pct(v, 1, True)}" for k, v in c["primes"].items()) + ").",
                "",
                tableau(["Budget", "Valeur médiane", "Fourchette 5 % – 95 %", "Probabilité de perte",
                         "Probabilité du plancher (pire hypothèse)", "Référence passive (médiane)",
                         "Fait mieux que la référence"],
                        [ligne_budget(float(b), r) for b, r in c["budgets"].items()]),
                "", "La dernière colonne compare à hypothèses égales : elle mesure l'effet du risque, des frais et des "
                    "stops, pas un avantage que l'analyse aurait sur le marché.", ""]
        r5 = c["budgets"].get("5000") or next(iter(c["budgets"].values()))
        if r5.get("rejeu"):
            rj = r5["rejeu"]
            out += [f"Rejeu des {c['fenetres_rejouees']} fenêtres réelles de {c['jours_bourse']} séances de l'historique "
                    f"(une tous les 5 jours, crises comprises, dans leur ordre) pour 5 000 € : médiane "
                    f"{eur(rj['mediane'])}, pire cas {eur(rj['pire_valeur'])}, perte dans {pct(rj['proba_perte'], 0)} "
                    f"des fenêtres, plancher atteint dans {pct(rj['proba_plancher'], 1)}.", ""]
        out += [f"Pour 5 000 €, la médiane vaut {eur(r5['neutre']['mediane'])} sans aucune prime de risque, et "
                f"{eur(r5['historique']['mediane'])} si chaque ligne prolongeait sa tendance passée (hypothèse "
                f"optimiste). Pire repli en cours de route (1 cas sur 20) : {pct(r5['central']['perte_max_p95'])}. "
                + (f"Imposition des gains : {pct(c['impot']['taux'])} ({c['impot']['motif']}), soit un gain médian net "
                   f"de {eur(r5['gain_median_net_impot'])} pour 5 000 €." if r5["gain_median_net_impot"] > 0 else
                   f"Perte médiane de {eur(-r5['gain_median_net_impot'])} pour 5 000 € : pas d'impôt, le scénario "
                   "central ne rémunère pas le risque pris."), ""]

    if backtest:
        out += ["## 7. Modèles testés hors échantillon", "",
                f"Fenêtres glissantes : paramètres choisis sur {backtest['entrainement_jours']} séances, testés sur les "
                f"{backtest['test_jours']} suivantes, frais compris (budget de 1 000 €, le plus pénalisé par les frais "
                f"fixes). Série : {backtest['fichier']}, arrêtée au {backtest['date_arrete']}.", ""]
        lig = []
        for nom, r in sorted(backtest["resultats"].items(), key=lambda kv: -(kv[1].get("sharpe") or -9)):
            lig.append([r["libelle"], ", ".join(r["horizons"]), pct(r["rendement_annuel"]), pct(r["volatilite"]),
                        nb(r["sharpe"]), pct(r["perte_max"]), pct(r["temps_investi"], 0),
                        nb(r["transactions_par_an"], 1), f"{nb(r['sharpe_entrainement'])} → {nb(r['sharpe'])}"])
        out += [tableau(["Modèle", "Horizons", "Rendement annuel", "Volatilité", "Sharpe", "Perte max", "Temps investi",
                         "Transactions / an", "Sharpe entraînement → test"], lig), ""]
        out += ["Un modèle dont le Sharpe s'effondre entre l'entraînement et le test est surajusté : il décrit le passé "
                "sans rien prédire.", ""]

    prevs = analyse.get("previsions") or []
    if prevs:
        out += ["## 8. Prévisions enregistrées au journal", "",
                tableau(["Id", "Énoncé", "Échéance", "Probabilité de l'agent", "Probabilité du modèle"],
                        [[p["id"], p.get("enonce"), p.get("echeance"), pct(p.get("probabilite"), 0),
                          pct(calc["previsions"].get(p["id"], {}).get("probabilite_modele"), 0)] for p in prevs]),
                "", "Chaque prévision sera notée à son échéance (score de Brier) contre la probabilité du modèle "
                    "statistique : c'est la mesure de ce que l'analyse apporte réellement.", ""]
    out += ["## 9. Pré-mortem : pourquoi cela pourrait échouer", ""]
    out += [f"- {x}" for x in an.get("premortem") or []] + [""]
    if analyse.get("suivi") or analyse.get("lecons_appliquees"):
        out += ["## 10. Suivi", ""]
        out += [f"- {x}" for x in analyse.get("suivi") or []]
        out += [f"- Leçon appliquée : {x}" for x in analyse.get("lecons_appliquees") or []]
        out.append("")
    out += ["## Sources des données et limites", ""]
    for f, info in sources.items():
        out.append(f"- {f} : {info.get('source')}, relevé le {info.get('recupere_le', 'n.d.')}"
                   + (" — SYNTHÉTIQUE" if info.get("synthetique") else ""))
    out += ["- Simulations : rééchantillonnage par blocs de l'historique ; elles ne contiennent aucun événement plus "
            "grave que ceux déjà vécus par la série, et un stop s'exécute au cours de clôture qui le franchit "
            "(un écart d'ouverture est donc compté, pas un écart en séance).",
            f"- Fiscalité : barèmes vérifiés le {cfg['fiscalite']['verifie_le']} ({cfg['fiscalite']['source']}).",
            "- Ce document aide Patrick à décider pour son propre compte. Il ne constitue ni un conseil en "
            "investissement personnalisé délivré à un tiers, ni une sollicitation. Les performances passées ne "
            "préjugent pas des performances futures ; une perte en capital est possible.", ""]
    return "\n".join(out)


INLINE_RE = re.compile(r"\*\*(.+?)\*\*|(?<![\w])_(.+?)_(?![\w])")


def inline_html(t):
    t = html.escape(t)
    return INLINE_RE.sub(lambda m: f"<strong>{m.group(1)}</strong>" if m.group(1) else f"<em>{m.group(2)}</em>", t)


def md_vers_html(texte):
    out, liste, i = [], False, 0
    lignes = texte.split("\n")
    while i < len(lignes):
        l = lignes[i]
        if not l.startswith("- ") and liste:
            out.append("</ul>")
            liste = False
        if l.startswith("|"):
            bloc = []
            while i < len(lignes) and lignes[i].startswith("|"):
                bloc.append([c.strip() for c in lignes[i].strip().strip("|").split("|")])
                i += 1
            tete, corps = bloc[0], [b for b in bloc[2:]]
            out.append("<table><thead><tr>" + "".join(f"<th>{inline_html(c)}</th>" for c in tete) + "</tr></thead><tbody>"
                       + "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>" for r in corps)
                       + "</tbody></table>")
            continue
        m = re.match(r"(#{1,3}) (.*)", l)
        if m:
            out.append(f"<h{len(m.group(1))}>{inline_html(m.group(2))}</h{len(m.group(1))}>")
        elif l.startswith("- "):
            if not liste:
                out.append("<ul>")
                liste = True
            out.append(f"<li>{inline_html(l[2:])}</li>")
        elif l.startswith("> "):
            out.append(f"<blockquote>{inline_html(l[2:])}</blockquote>")
        elif l.strip():
            out.append(f"<p>{inline_html(l)}</p>")
        i += 1
    if liste:
        out.append("</ul>")
    return "\n".join(out)


def docx_secours(texte, chemin):
    """DOCX par python-docx quand pandoc manque ; renvoie False si le module manque aussi."""
    try:
        from docx import Document
    except ImportError:
        return False
    doc = Document()
    lignes, i = texte.split("\n"), 0

    def remplir(p, t):
        pos = 0
        for m in INLINE_RE.finditer(t):
            p.add_run(t[pos:m.start()])
            r = p.add_run(m.group(1) or m.group(2))
            r.bold, r.italic = bool(m.group(1)), bool(m.group(2))
            pos = m.end()
        p.add_run(t[pos:])

    while i < len(lignes):
        l = lignes[i]
        if l.startswith("|"):
            bloc = []
            while i < len(lignes) and lignes[i].startswith("|"):
                bloc.append([c.strip() for c in lignes[i].strip().strip("|").split("|")])
                i += 1
            rangs = [bloc[0]] + bloc[2:]
            tb = doc.add_table(rows=len(rangs), cols=len(bloc[0]))
            tb.style = "Table Grid"
            for j, r in enumerate(rangs):
                for k, c in enumerate(r[:len(bloc[0])]):
                    remplir(tb.rows[j].cells[k].paragraphs[0], f"**{c}**" if j == 0 and c else c)
            continue
        m = re.match(r"(#{1,3}) (.*)", l)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif l.startswith("- "):
            remplir(doc.add_paragraph(style="List Bullet"), l[2:])
        elif l.startswith("> "):
            remplir(doc.add_paragraph(style="Intense Quote"), l[2:])
        elif l.strip():
            remplir(doc.add_paragraph(), l)
        i += 1
    doc.save(str(chemin))
    return True


def ecrire_livrables(sortie, base, texte, titre, R):
    sortie.mkdir(parents=True, exist_ok=True)
    (sortie / f"{base}.md").write_text(texte, encoding="utf-8")
    page = (f'<!doctype html><html lang="fr"><head><meta charset="utf-8"><title>{html.escape(titre)}</title>'
            f'<meta name="viewport" content="width=device-width, initial-scale=1"><style>{CSS}</style></head>'
            f'<body><div class="feuille">{md_vers_html(texte)}</div></body></html>')
    (sortie / f"{base}.html").write_text(page, encoding="utf-8")
    if shutil.which("pandoc"):
        subprocess.run(["pandoc", "-f", "gfm", "-t", "docx", "-o", str(sortie / f"{base}.docx")],
                       input=texte, text=True, check=True)
    elif not docx_secours(texte, sortie / f"{base}.docx"):
        R.alerte("ni pandoc ni python-docx : DOCX non produit")
    for ext in ("md", "html", "docx"):
        if (sortie / f"{base}.{ext}").exists():
            print(f"  {sortie / base}.{ext}")


# ── Journal ──────────────────────────────────────────────────────────────

def lire_journal(dossier_journal):
    f = Path(dossier_journal) / "previsions.jsonl"
    if not f.exists():
        return []
    return [json.loads(l) for l in f.read_text(encoding="utf-8").splitlines() if l.strip()]


def ecrire_journal(dossier_journal, entrees):
    d = Path(dossier_journal)
    d.mkdir(parents=True, exist_ok=True)
    (d / "previsions.jsonl").write_text(
        "".join(json.dumps(e, ensure_ascii=False, default=_json_defaut) + "\n" for e in entrees), encoding="utf-8")


VERROUILLES = ("probabilite", "condition", "echeance", "fichier")


def memoriser_initiales(dossier, analyse):
    """Garde la première probabilité écrite pour chaque prévision, avant que l'agent voie celle du modèle."""
    chemin = dossier / "previsions-initiales.json"
    initiales = json.loads(chemin.read_text(encoding="utf-8")) if chemin.exists() else {}
    nouveau = False
    for p in analyse.get("previsions") or []:
        if p.get("id") and p["id"] not in initiales and fini(p.get("probabilite")):
            initiales[p["id"]] = {"probabilite": p["probabilite"],
                                  "le": dt.datetime.now().isoformat(timespec="seconds")}
            nouveau = True
    if nouveau:
        ecrire_json(chemin, initiales)
    return initiales


def journaliser(dossier, analyse, calc, journal, R, initiales):
    entrees = lire_journal(journal)
    par_id = {e["id"]: i for i, e in enumerate(entrees)}
    ref = analyse.get("reference_passive") or {}
    meta = analyse["meta"]
    ajoutees = 0
    for p in analyse.get("previsions") or []:
        ident = f"{meta['reference']}/{p['id']}"
        c = calc["previsions"].get(p["id"], {})
        entree = {"id": ident, "reference": meta["reference"], "cree_le": calc["date_arrete"],
                  "dossier": str(dossier.resolve()), "horizon": p.get("horizon"), "enonce": p.get("enonce"),
                  "instrument": analyse["instrument"].get("nom"), "type_instrument": analyse["instrument"].get("type"),
                  "symbole": p.get("symbole") or analyse["instrument"].get("symbole"), "fichier": p.get("fichier"),
                  "echeance": p.get("echeance"), "condition": p.get("condition"), "probabilite": p.get("probabilite"),
                  "probabilite_modele": c.get("probabilite_modele"), "prix_reference": c.get("prix_reference"),
                  "reference_passive": {"fichier": ref.get("fichier"), "nom": ref.get("nom")},
                  "strategie": p.get("strategie"), "lecons_appliquees": analyse.get("lecons_appliquees", []),
                  "probabilite_initiale": initiales.get(p["id"], {}).get("probabilite"), "statut": "ouverte"}
        if ident in par_id:
            ancien = entrees[par_id[ident]]
            if ancien.get("statut") != "ouverte":
                R.alerte(f"prévision {ident} déjà évaluée : non remplacée")
                continue
            modifies = [k for k in VERROUILLES if ancien.get(k) != entree[k]]
            if modifies:
                R.erreur(f"prévision {ident} inscrite le {ancien.get('cree_le')} : {', '.join(modifies)} ne se "
                         "modifie plus, le score serait faussé ; créer une nouvelle prévision (nouvel identifiant)")
                continue
            # Les valeurs du jour d'inscription restent celles qui seront notées.
            entree.update({k: ancien[k] for k in ("cree_le", "probabilite_modele", "prix_reference",
                                                    "probabilite_initiale") if k in ancien})
            entrees[par_id[ident]] = entree
        else:
            entrees.append(entree)
            ajoutees += 1
    ecrire_journal(journal, entrees)
    print(f"  journal : {Path(journal) / 'previsions.jsonl'} ({ajoutees} prévision(s) ajoutée(s), "
          f"{len(entrees)} au total)")


def chercher_fichier(rel, dossiers, symbole=None):
    """Série d'une prévision : d'abord dans le dossier d'analyse, puis dans les dossiers fournis (--donnees)."""
    if not rel and not symbole:
        return None
    nom = Path(rel or "").name
    candidats = []
    for d in filter(None, dossiers):
        d = Path(d)
        if rel:
            candidats += [d / rel, d / nom, d / "donnees" / nom]
        if symbole:
            candidats += [d / f"{slug(symbole)}.csv", d / "donnees" / f"{slug(symbole)}.csv"]
    return next((c for c in candidats if c.is_file()), None)


def cmd_evaluer(a):
    jour = aujourdhui(a)
    entrees = lire_journal(a.journal)
    if not entrees:
        sortir(f"journal vide : {Path(a.journal) / 'previsions.jsonl'}")
    imposes = {}
    for x in a.prix or []:
        sym, _, v = x.partition("=")
        if nombre(v) is None:
            sortir(f"--prix {x} : SYMBOLE=VALEUR attendu")
        imposes[sym.strip()] = nombre(v)
    echues = [e for e in entrees if e.get("statut") == "ouverte" and dt.date.fromisoformat(e["echeance"]) <= jour]
    for cle in imposes:
        if cle not in {e["id"] for e in echues} and sum(e.get("symbole") == cle for e in echues) > 1:
            sortir(f"--prix {cle} : plusieurs prévisions échues sur ce symbole, à des dates différentes ; "
                   "donner l'identifiant de la prévision (--prix <id>=VALEUR)")
    evaluees, manquantes = 0, []
    for e in echues:
        ech = dt.date.fromisoformat(e["echeance"])
        cle = e["id"] if e["id"] in imposes else e.get("symbole")
        if cle in imposes:
            prix, date_prix = imposes[cle], ech
        else:
            f = chercher_fichier(e.get("fichier"), [e.get("dossier")] + (a.donnees or []), e.get("symbole"))
            if not f:
                manquantes.append(f"{e['id']} ({e.get('symbole') or e.get('fichier')}, échéance {ech})")
                continue
            dates, p = lire_serie(f)
            date_prix, prix = cours_a(dates, p, ech)
            if date_prix is None or (ech - date_prix).days > 5:
                manquantes.append(f"{e['id']} : la série {f.name} s'arrête avant l'échéance {ech}")
                continue
        cond = e["condition"]
        issue = 1 if (prix > cond["seuil"] if cond["type"] == "au_dessus" else prix < cond["seuil"]) else 0
        e.update({"statut": "evaluee", "evaluee_le": jour.isoformat(), "prix_realise": prix,
                  "date_prix_realise": date_prix, "issue": issue,
                  "brier_agent": (e["probabilite"] - issue) ** 2,
                  "brier_agent_initial": (e["probabilite_initiale"] - issue) ** 2
                  if e.get("probabilite_initiale") is not None else None,
                  "brier_modele": (e["probabilite_modele"] - issue) ** 2 if e.get("probabilite_modele") is not None else None,
                  "rendement_realise": prix / e["prix_reference"] - 1 if e.get("prix_reference") else None})
        rp = e.get("reference_passive") or {}
        fr = chercher_fichier(rp.get("fichier"), [e.get("dossier")] + (a.donnees or []))
        if fr:
            dr, pr = lire_serie(fr)
            _, v0 = cours_a(dr, pr, dt.date.fromisoformat(str(e["cree_le"])))
            _, v1 = cours_a(dr, pr, ech)
            if v0 and v1:
                e["rendement_reference"] = v1 / v0 - 1
        evaluees += 1
        print(f"  {e['id']} : {'réalisée' if issue else 'non réalisée'} (cours {nb(prix)} le {date_prix}, "
              f"agent {pct(e['probabilite'], 0)}, modèle {pct(e.get('probabilite_modele'), 0)})")
    ecrire_journal(a.journal, entrees)
    print(f"\n{evaluees} prévision(s) évaluée(s).")
    if manquantes:
        print("Cours manquants, à fournir par --prix SYMBOLE=VALEUR ou --donnees DOSSIER :")
        for m in manquantes:
            print(f"  - {m}")
    a_annoter = [e["id"] for e in entrees if e.get("statut") == "evaluee" and not e.get("issue") and not e.get("cause")]
    if a_annoter:
        print("Prévisions manquées à annoter (finance.py annoter <id> --cause ...) : " + ", ".join(a_annoter))
    return 0


def cmd_annoter(a):
    entrees = lire_journal(a.journal)
    for e in entrees:
        if e["id"] == a.id:
            e["cause"] = a.cause
            if a.note:
                e["note"] = a.note
            ecrire_journal(a.journal, entrees)
            print(f"{a.id} : cause « {a.cause} » enregistrée")
            return 0
    sortir(f"prévision {a.id} absente du journal")


def brier(es, cle):
    v = [e[cle] for e in es if e.get(cle) is not None]
    return sum(v) / len(v) if v else None


def cmd_bilan(a):
    jour = aujourdhui(a)
    entrees = lire_journal(a.journal)
    ev = [e for e in entrees if e.get("statut") == "evaluee"]
    ouvertes = [e for e in entrees if e.get("statut") == "ouverte"]
    en_retard = [e for e in ouvertes if dt.date.fromisoformat(e["echeance"]) <= jour]
    out = [f"# Bilan de l'expert financier au {jour}", "",
           f"{len(entrees)} prévision(s) au journal : {len(ev)} évaluée(s), {len(ouvertes)} ouverte(s) "
           f"dont {len(en_retard)} échue(s) à évaluer.", ""]
    signaux = []
    if ev:
        ba, bm = brier(ev, "brier_agent"), brier([e for e in ev if e.get("brier_modele") is not None], "brier_modele")
        comp = [e for e in ev if e.get("brier_modele") is not None]
        ba_c = brier(comp, "brier_agent")
        habilete = 1 - ba_c / bm if bm and ba_c is not None and bm > 0 else None
        ini = [e for e in comp if e.get("brier_agent_initial") is not None]
        out += ["## Qualité des prévisions", "",
                tableau(["Mesure", "Valeur"], [
                    ["Prévisions évaluées", str(len(ev))],
                    ["Taux de réussite", pct(sum(e["issue"] for e in ev) / len(ev), 0)],
                    ["Score de Brier de l'agent (0 = parfait, 0,25 = pile ou face)", nb(ba, 3)],
                    ["Score de Brier du modèle statistique", nb(bm, 3)],
                    ["Apport de l'analyse (1 − Brier agent / Brier modèle)", pct(habilete, 0, True)]]
                    + ([["Brier de l'agent avant d'avoir vu le modèle", nb(brier(ini, "brier_agent_initial"), 3)],
                        ["Apport du jugement seul, avant ancrage sur le modèle",
                         pct(1 - brier(ini, "brier_agent_initial") / brier(ini, "brier_modele"), 0, True)
                         if brier(ini, "brier_modele") else "n.d."]] if ini else [])), ""]
        tranches = [(0, .2), (.2, .4), (.4, .6), (.6, .8), (.8, 1.01)]
        lig = []
        for lo, hi in tranches:
            g = [e for e in ev if lo <= e["probabilite"] < hi]
            if g:
                pm = sum(e["probabilite"] for e in g) / len(g)
                fr = sum(e["issue"] for e in g) / len(g)
                lig.append([f"{pct(lo, 0)} – {pct(min(hi, 1), 0)}", str(len(g)), pct(pm, 0), pct(fr, 0)])
                if len(g) >= 5 and pm >= 0.6 and pm - fr > 0.10:
                    signaux.append(f"Surconfiance : annoncé {pct(pm, 0)}, réalisé {pct(fr, 0)} sur {len(g)} prévisions.")
                if len(g) >= 5 and pm <= 0.4 and fr - pm > 0.10:
                    signaux.append(f"Surconfiance à la baisse : annoncé {pct(pm, 0)}, réalisé {pct(fr, 0)} "
                                   f"sur {len(g)} prévisions.")
        out += ["### Calibrage", "", tableau(["Probabilité annoncée", "Nombre", "Moyenne annoncée", "Fréquence réalisée"],
                                             lig), ""]
        lig = []
        for cle, lib in (("horizon", "Horizon"), ("type_instrument", "Type d'instrument")):
            for v in sorted({str(e.get(cle)) for e in ev}):
                g = [e for e in ev if str(e.get(cle)) == v]
                gc = [e for e in g if e.get("brier_modele") is not None]
                h_ = 1 - brier(gc, "brier_agent") / brier(gc, "brier_modele") if gc and brier(gc, "brier_modele") else None
                lig.append([lib, v, str(len(g)), nb(brier(g, "brier_agent"), 3), pct(h_, 0, True)])
                if h_ is not None and len(g) >= 10 and h_ < 0:
                    signaux.append(f"{lib} « {v} » : l'analyse fait moins bien que le modèle sur {len(g)} prévisions.")
        out += ["### Par horizon et par type", "",
                tableau(["Axe", "Valeur", "Nombre", "Brier agent", "Apport"], lig), ""]
        alphas = [e["rendement_realise"] - e["rendement_reference"] for e in ev
                  if e.get("rendement_realise") is not None and e.get("rendement_reference") is not None]
        if alphas:
            out += [f"Écart moyen de l'instrument choisi sur la référence passive, entre la prévision et l'échéance : "
                    f"{pct(sum(alphas) / len(alphas), 1, True)} sur {len(alphas)} prévision(s).", ""]
        manquees = [e for e in ev if not e["issue"]]
        if manquees:
            causes = {c: sum(1 for e in manquees if e.get("cause") == c) for c in CAUSES}
            causes["non annotée"] = sum(1 for e in manquees if not e.get("cause"))
            out += ["### Causes des prévisions manquées", "",
                    tableau(["Cause", "Nombre"], [[c, str(n)] for c, n in causes.items() if n]), ""]
            if causes["evenement_exogene"] >= max(3, 0.3 * len(manquees)):
                signaux.append("Un tiers au moins des échecs vient d'un événement exogène : élargir la veille et "
                               "prévoir les parades avant l'événement.")
            if causes["non annotée"]:
                signaux.append(f"{causes['non annotée']} échec(s) sans cause : les annoter avant d'en tirer une leçon.")
        if len(ev) < 20:
            signaux.insert(0, f"Échantillon de {len(ev)} prévisions : trop petit pour conclure (20 au moins). "
                              "Aucune leçon statistique à en tirer pour l'instant.")
        elif habilete is not None and habilete < 0:
            signaux.append("L'analyse fait moins bien que le modèle statistique : s'en rapprocher, et ne s'en écarter "
                           "qu'avec un argument nouveau, précis et sourcé.")
    registre = Path(a.journal) / "registre-modeles.jsonl"
    if registre.exists():
        lignes = [json.loads(l) for l in registre.read_text(encoding="utf-8").splitlines() if l.strip()]
        par = {}
        for l in lignes:
            par.setdefault(l["strategie"], []).append(l)
        lig = []
        for nom, g in sorted(par.items()):
            gains = sorted(x["sharpe"] - x["sharpe_achat"] for x in g if x.get("sharpe") is not None
                           and x.get("sharpe_achat") is not None)
            pertes = sorted(x["perte_max"] - x["perte_max_achat"] for x in g if x.get("perte_max_achat") is not None)
            instr = len({x["fichier"] for x in g})
            med = gains[len(gains) // 2] if gains else None
            medp = pertes[len(pertes) // 2] if pertes else None
            candidat = instr >= 3 and med is not None and med > 0.1 and (medp is None or medp >= -0.05)
            lig.append([g[0].get("libelle", nom), str(instr), nb(med), pct(medp, 1, True), "oui" if candidat else "non"])
        out += ["## Registre des modèles", "",
                tableau(["Modèle", "Instruments testés", "Gain de Sharpe médian sur l'achat-conservation",
                         "Écart de perte max médian", "Candidat champion"], lig),
                "", "Candidat champion : testé sur trois instruments au moins, gain de Sharpe médian supérieur à 0,1, "
                    "perte maximale pas plus de 5 points plus profonde.", ""]
    lecons = Path(a.journal) / "lecons.md"
    if lecons.exists():
        vieilles = []
        for l in lecons.read_text(encoding="utf-8").splitlines():
            m = re.match(r"- \[(\d{4}-\d{2}-\d{2})\]\s*(.*)", l)
            if m and (jour - dt.date.fromisoformat(m.group(1))).days > 180:
                vieilles.append(f"{m.group(1)} : {m.group(2)[:90]}")
        if vieilles:
            signaux.append("Leçons de plus de six mois à réexaminer : " + " ; ".join(vieilles))
    out += ["## Signaux pour l'auto-amélioration", ""]
    out += [f"- {s}" for s in signaux] or ["- Aucun signal."]
    out += ["", "Une leçon n'entre dans lecons.md qu'après accord de Patrick, avec sa date et sa preuve. "
                "Les règles de risque (plancher, produits interdits) ne changent jamais par ce biais.", ""]
    if en_retard:
        out += ["## Prévisions échues à évaluer", ""] + [f"- {e['id']} : {e['enonce']} (échéance {e['echeance']})"
                                                         for e in en_retard] + [""]
    texte = "\n".join(out)
    sortie = Path(a.sortie or a.journal)
    ecrire_livrables(sortie, f"bilan-{jour}", texte, f"Bilan {jour}", Rapport())
    print("\n" + "\n".join(f"  {s}" for s in signaux))
    return 0


# ── Commandes ────────────────────────────────────────────────────────────

def exiger_ignore(dossier, racine):
    r = subprocess.run(["git", "check-ignore", "-q", str(dossier)], capture_output=True)
    if r.returncode == 1:
        shutil.rmtree(dossier)
        sortir(f"{dossier} n'est pas ignoré par git : ajouter « {racine}/ » au .gitignore (dépôt public)")


def cmd_init(a):
    racine = Path(a.racine)
    date = a.date or dt.date.today().isoformat()
    ref = f"{date}-{slug(a.instrument)}"
    dossier = racine / "analyses" / ref
    if dossier.exists():
        sortir(f"{dossier} existe déjà")
    (dossier / "donnees").mkdir(parents=True)
    exiger_ignore(dossier, racine)
    (racine / "journal").mkdir(exist_ok=True)
    (racine / "modeles").mkdir(exist_ok=True)
    squelette = {
        "meta": {"reference": ref, "date": date, "demande": "", "statut": "brouillon"},
        "parametres": {"enveloppe": "CTO"},
        "instrument": {"nom": a.instrument, "type": "", "symbole": "", "isin": "", "place": "", "devise": "EUR",
                       "eligible_pea": None, "fichier": "", "prix": None, "prix_date": "", "prix_source": ""},
        "reference_passive": {"nom": "", "fichier": ""},
        "analyse": {"these": "", "fondamental": [], "valorisation": [], "technique": "", "premortem": []},
        "evenements": [], "scenarios": [], "strategies": {}, "previsions": [], "recommandation": "", "suivi": [],
        "lecons_appliquees": [],
    }
    ecrire_json(dossier / "analyse.json", squelette)
    (dossier / "journal.md").write_text(f"# Journal — {ref}\n\n", encoding="utf-8")
    print(dossier)
    return 0


def cmd_donnees(a):
    dossier = Path(a.dossier)
    debut = dt.date.fromisoformat(a.debut)
    points, source, ajuste = telecharger(a.source, a.symbole, debut)
    nom = a.nom or slug(a.symbole)
    ecrire_serie(dossier / "donnees" / f"{nom}.csv", points)
    noter_source(dossier, nom, {"source": source, "symbole": a.symbole, "recupere_le": dt.date.today().isoformat(),
                                "debut": points[0][0], "fin": points[-1][0], "ajuste_dividendes": ajuste})
    print(f"{dossier / 'donnees' / nom}.csv : {len(points)} cours du {points[0][0]} au {points[-1][0]}")
    return 0


def cmd_importer(a):
    dossier = Path(a.dossier)
    points, colonne = normaliser_csv(Path(a.fichier).read_text(encoding="utf-8-sig", errors="replace"), a.colonne)
    ajuste = True if re.search(r"adj|ajust", sans_accent(colonne)) else a.ajuste
    ecrire_serie(dossier / "donnees" / f"{a.nom}.csv", points)
    noter_source(dossier, a.nom, {"source": a.source or f"import de {Path(a.fichier).name}",
                                  "recupere_le": dt.date.today().isoformat(), "debut": points[0][0],
                                  "fin": points[-1][0], "colonne": colonne, "ajuste_dividendes": ajuste})
    if not ajuste:
        print(f"Colonne retenue : « {colonne} ». Cours non ajustés des dividendes, sauf indication contraire "
              "(--ajuste si la source les ajuste).")
    print(f"{dossier / 'donnees' / a.nom}.csv : {len(points)} cours du {points[0][0]} au {points[-1][0]}")
    return 0


def cmd_analyser(a):
    dossier = Path(a.dossier)
    asof = aujourdhui(a)
    analyse = json.loads((dossier / "analyse.json").read_text(encoding="utf-8")) \
        if (dossier / "analyse.json").exists() else {}
    cfg = charger_config(analyse)
    rf = cfg["taux_sans_risque"]["valeur"]
    ref_f = (analyse.get("reference_passive") or {}).get("fichier")
    ref = lire_serie(dossier / ref_f, asof) if ref_f and (dossier / ref_f).exists() else None
    resultats = {}
    for f in sorted((dossier / "donnees").glob("*.csv")):
        resultats[f"donnees/{f.name}"] = analyser_serie(*lire_serie(f, asof), rf,
                                                        ref if f"donnees/{f.name}" != ref_f else None)
    ecrire_json(dossier / "analyse-quant.json", resultats)
    for f, q in resultats.items():
        print(f"\n{f} ({q['debut']} → {q['fin']}, dernier cours {nb(q['dernier_cours'])})")
        print(f"  rendement annualisé {pct(q.get('rendement_annuel'))}, volatilité {pct(q.get('volatilite'))}, "
              f"Sharpe {nb(q.get('sharpe'))}, perte max {pct(q['perte_max_detail']['perte_max'])}")
        print(f"  1 an {pct(q['performances']['1_an'])}, momentum 12-1 {pct(q.get('momentum_12_1'))}, "
              f"RSI14 {nb(q.get('rsi14'), 0)}, MM200 {nb(q.get('mm200'))}, vol 20 j {pct(q.get('volatilite_20j'))}")
    print(f"\n{dossier / 'analyse-quant.json'}")
    return 0


def cmd_backtest(a):
    dossier = Path(a.dossier)
    asof = aujourdhui(a)
    analyse = json.loads((dossier / "analyse.json").read_text(encoding="utf-8")) \
        if (dossier / "analyse.json").exists() else {}
    cfg = charger_config(analyse)
    R = Rapport()
    rel = a.fichier or (analyse.get("instrument") or {}).get("fichier")
    if not rel:
        sortir("préciser --fichier (aucun instrument.fichier dans analyse.json)")
    dates, p = lire_serie(dossier / rel, asof)
    rf = cfg["taux_sans_risque"]["valeur"]
    bt = cfg["backtest"]
    if len(p) < bt["entrainement_jours"] + bt["test_jours"] + 1:
        sortir(f"{rel} : historique trop court pour un test hors échantillon "
               f"({bt['entrainement_jours'] + bt['test_jours']} séances au moins)")
    strategies = dict(STRATEGIES)
    strategies.update(charger_modeles([SKILL / "modeles", Path("finance/modeles")] + (a.modeles or []), R))
    if a.signaux:
        pts = dict(normaliser_csv(Path(a.signaux).read_text(encoding="utf-8-sig"), a.colonne_signaux,
                                  positif=False)[0])
        e = np.nan_to_num(np.array([pts.get(d, np.nan) for d in dates]))
        strategies["signaux_externes"] = {"fn": None, "grille": [{}], "horizons": ["?"],
                                          "libelle": f"Signaux importés ({Path(a.signaux).name})", "expo": e}
        R.alerte("signaux importés : l'absence d'anticipation ne peut pas être contrôlée, vérifier comment ils ont "
                 "été produits (chaque signal ne doit utiliser que des données antérieures)")
    budget_min = min(cfg["budgets"])
    resultats, registre = {}, []
    achat = walk_forward(p, STRATEGIES["achat_conservation"], rf, cfg["courtier"], budget_min,
                         bt["entrainement_jours"], bt["test_jours"])
    for nom, s in strategies.items():
        if s.get("expo") is not None:
            s = dict(s, fn=lambda prix, _e=s["expo"]: _e[:len(prix)])
        elif not all(sans_anticipation(s["fn"], p, g) for g in s["grille"]):
            R.erreur(f"modèle {nom} : l'exposition d'un jour dépend de cours postérieurs (anticipation), écarté")
            continue
        res = walk_forward(p, s, rf, cfg["courtier"], budget_min, bt["entrainement_jours"], bt["test_jours"])
        if not res:
            continue
        res.update({"libelle": s["libelle"], "horizons": s["horizons"]})
        par_budget = {}
        for b in cfg["budgets"]:
            r_b = walk_forward(p, s, rf, cfg["courtier"], b, bt["entrainement_jours"], bt["test_jours"])
            par_budget[str(b)] = r_b["rendement_annuel"]
        res["rendement_par_budget"] = par_budget
        resultats[nom] = res
        if res["sharpe"] is not None and res["sharpe_entrainement"] - res["sharpe"] > 0.5:
            R.alerte(f"modèle {nom} : Sharpe {nb(res['sharpe_entrainement'])} à l'entraînement, "
                     f"{nb(res['sharpe'])} en test : surajustement probable")
        registre.append({"date": dt.date.today().isoformat(), "date_arrete": asof, "fichier": rel,
                         "instrument": (analyse.get("instrument") or {}).get("nom"), "strategie": nom,
                         "libelle": s["libelle"], "horizons": s["horizons"], "sharpe": res["sharpe"],
                         "rendement_annuel": res["rendement_annuel"], "perte_max": res["perte_max"],
                         "sharpe_achat": achat["sharpe"], "perte_max_achat": achat["perte_max"],
                         "parametres": res["parametres_retenus"]})
    sortie = {"date_arrete": asof, "fichier": rel, "entrainement_jours": bt["entrainement_jours"],
              "test_jours": bt["test_jours"], "budget": budget_min, "resultats": resultats}
    ecrire_json(dossier / "backtest.json", sortie)
    journal = Path(a.journal)
    journal.mkdir(parents=True, exist_ok=True)
    chemin = journal / "registre-modeles.jsonl"
    cle = lambda l: (l.get("instrument"), l["fichier"], str(l["date_arrete"]), l["strategie"])
    anciens = [json.loads(l) for l in chemin.read_text(encoding="utf-8").splitlines() if l.strip()] \
        if chemin.exists() else []
    nouveaux = {cle(l) for l in registre}
    garde = [l for l in anciens if cle(l) not in nouveaux] + registre
    chemin.write_text("".join(json.dumps(l, ensure_ascii=False, default=_json_defaut) + "\n" for l in garde),
                      encoding="utf-8")
    print(f"Test hors échantillon de {rel} ({dates[0]} → {dates[-1]}), frais d'un budget de {eur(budget_min)} :")
    print(f"  {'modèle':<22}{'rdt/an':>9}{'vol':>8}{'Sharpe':>8}{'perte max':>11}{'investi':>9}{'trans/an':>9}"
          f"{'Sh. entr.':>10}")
    for nom, r in sorted(resultats.items(), key=lambda kv: -(kv[1]["sharpe"] or -9)):
        print(f"  {nom:<22}{pct(r['rendement_annuel']):>9}{pct(r['volatilite']):>8}{nb(r['sharpe']):>8}"
              f"{pct(r['perte_max']):>11}{pct(r['temps_investi'], 0):>9}{nb(r['transactions_par_an'], 1):>9}"
              f"{nb(r['sharpe_entrainement']):>10}")
    print(f"\n  {dossier / 'backtest.json'}\n  {journal / 'registre-modeles.jsonl'}")
    R.afficher()
    return 1 if R.erreurs else 0


def cmd_construire(a):
    dossier = Path(a.dossier)
    analyse = lire_json(dossier / "analyse.json")
    asof = aujourdhui(a)
    cfg = charger_config(analyse)
    R = Rapport()
    initiales = memoriser_initiales(dossier, analyse)
    calc, _ = calculer(dossier, analyse, cfg, asof, R, a.journal, a.chemins)
    sources_f = dossier / "donnees" / "sources.json"
    sources = json.loads(sources_f.read_text(encoding="utf-8")) if sources_f.exists() else {}
    for f in {analyse.get("instrument", {}).get("fichier"), (analyse.get("reference_passive") or {}).get("fichier")}:
        if f and f not in sources:
            R.alerte(f"origine de {f} non documentée (donnees/sources.json)")
    actions = {l.get("proxy") or l.get("fichier") for s in (analyse.get("strategies") or {}).values()
               for l in (s or {}).get("lignes", []) if classe_de(l, cfg) == "actions"}
    actions |= {(analyse.get("reference_passive") or {}).get("fichier")}
    for f in sorted(x for x in actions if x and x in sources):
        if not sources[f].get("synthetique") and sources[f].get("ajuste_dividendes") is not True:
            R.alerte(f"{f} : cours non ajustés des dividendes (ou inconnu) ; l'historique, le rejeu et l'hypothèse "
                     "historique sous-estiment le rendement, pas le cas central")
    bt_f = dossier / "backtest.json"
    backtest = json.loads(bt_f.read_text(encoding="utf-8")) if bt_f.exists() else None
    if backtest and backtest.get("date_arrete") != asof.isoformat():
        R.alerte(f"backtest.json arrêté au {backtest.get('date_arrete')}, analyse au {asof} : relancer le backtest")
    if not backtest:
        R.alerte("aucun test hors échantillon (finance.py backtest) : les stratégies ne s'appuient sur aucune preuve")
    ecrire_json(dossier / "calculs.json", calc)
    texte = note_md(analyse, calc, cfg, R, backtest, sources)
    ref = analyse.get("meta", {}).get("reference", dossier.name)
    print(f"Note {ref} :")
    ecrire_livrables(dossier / "livrables", f"{ref}-note", texte, f"Note d'analyse {ref}", R)
    if a.final:
        if R.erreurs:
            print("\nNote non finale : erreurs à corriger, rien n'est inscrit au journal.")
        else:
            journaliser(dossier, analyse, calc, a.journal, R, initiales)
    R.afficher()
    return 1 if R.erreurs else 0


def cmd_obligation(a):
    jour = aujourdhui(a)
    o = obligation(a.prix, a.coupon, dt.date.fromisoformat(a.echeance), jour, a.frequence)
    print(f"Obligation coupon {nb(a.coupon)} %, échéance {a.echeance}, prix {nb(a.prix)} au {jour} :")
    for cle, v in o.items():
        print(f"  {cle:<30} {pct(v, 2) if 'rendement' in cle or 'variation' in cle else nb(v, 3)}")
    return 0


def cmd_archiver(a):
    racine = Path(a.racine).resolve()
    if not racine.is_dir():
        sortir(f"{racine} introuvable")
    # L'archive reste dans la racine, ignorée par git : jamais à côté, dans le dépôt public.
    cible = racine / f"sauvegarde-{dt.date.today().isoformat()}.zip"
    with zipfile.ZipFile(cible, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(racine.rglob("*")):
            if f.is_file() and f.suffix != ".zip" and "livrables" not in f.parts:
                z.write(f, Path(racine.name) / f.relative_to(racine))
    print(cible)
    return 0


def cmd_restaurer(a):
    racine = Path(a.racine)
    with zipfile.ZipFile(a.archive) as z:
        noms = z.namelist()
        if any(n.startswith("/") or ".." in Path(n).parts for n in noms):
            sortir("archive inattendue : chemins absolus ou remontants")
        tetes = {Path(n).parts[0] for n in noms if n}
        if len(tetes) != 1:
            sortir("archive inattendue : un seul dossier racine attendu")
        racine.mkdir(parents=True, exist_ok=True)
        exiger_ignore(racine, racine)
        for n in noms:
            rel = Path(*Path(n).parts[1:])
            if not rel.parts or n.endswith("/"):
                continue
            cible = racine / rel
            if cible.exists():
                print(f"  conservé (déjà présent) : {cible}")
                continue
            cible.parent.mkdir(parents=True, exist_ok=True)
            cible.write_bytes(z.read(n))
    print(racine)
    return 0


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = p.add_subparsers(dest="cmd", required=True)
    x = sp.add_parser("init")
    x.add_argument("racine", nargs="?", default="finance")
    x.add_argument("--instrument", required=True)
    x.add_argument("--date")
    x = sp.add_parser("donnees")
    x.add_argument("dossier")
    x.add_argument("--symbole", required=True)
    x.add_argument("--source", default="yahoo", choices=["yahoo", "stooq", "fred"])
    x.add_argument("--nom")
    x.add_argument("--debut", default="2000-01-01")
    x = sp.add_parser("importer")
    x.add_argument("dossier")
    x.add_argument("fichier")
    x.add_argument("--nom", required=True)
    x.add_argument("--colonne")
    x.add_argument("--source")
    x.add_argument("--ajuste", action="store_const", const=True, default=None,
                   help="la source fournit des cours ajustés des dividendes")
    for nom in ("analyser", "backtest", "construire"):
        x = sp.add_parser(nom)
        x.add_argument("dossier")
        x.add_argument("--aujourdhui")
        x.add_argument("--journal", default="finance/journal")
        if nom == "backtest":
            x.add_argument("--fichier")
            x.add_argument("--modeles", nargs="*")
            x.add_argument("--signaux")
            x.add_argument("--colonne-signaux", dest="colonne_signaux")
        if nom == "construire":
            x.add_argument("--final", action="store_true")
            x.add_argument("--chemins", type=int)
    x = sp.add_parser("obligation")
    x.add_argument("--prix", type=float, required=True)
    x.add_argument("--coupon", type=float, required=True)
    x.add_argument("--echeance", required=True)
    x.add_argument("--frequence", type=int, default=1, choices=[1, 2, 4, 12])
    x.add_argument("--aujourdhui")
    x = sp.add_parser("evaluer")
    x.add_argument("--journal", default="finance/journal")
    x.add_argument("--donnees", nargs="*")
    x.add_argument("--prix", nargs="*")
    x.add_argument("--aujourdhui")
    x = sp.add_parser("annoter")
    x.add_argument("id")
    x.add_argument("--cause", required=True, choices=CAUSES)
    x.add_argument("--note")
    x.add_argument("--journal", default="finance/journal")
    x = sp.add_parser("bilan")
    x.add_argument("--journal", default="finance/journal")
    x.add_argument("--sortie")
    x.add_argument("--aujourdhui")
    x = sp.add_parser("archiver")
    x.add_argument("racine", nargs="?", default="finance")
    x = sp.add_parser("restaurer")
    x.add_argument("archive")
    x.add_argument("racine", nargs="?", default="finance")
    a = p.parse_args()
    return globals()[f"cmd_{a.cmd}"](a)


if __name__ == "__main__":
    sys.exit(main())
